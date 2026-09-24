"""
Internal Audit Observations Dashboard
=====================================
A Streamlit dashboard for recording, classifying and reporting internal audit
observations to an Audit Committee.

What it does
------------
* Reads one or more Excel/CSV observation schedules (header rows are detected
  automatically, column names are matched by synonym, missing fields are fine).
* Auto-classifies every observation into an internal-control area and a COSO
  component, detects cross-cutting themes, and labels design vs operating gaps.
* Auto-suggests a risk rating (likelihood x impact) wherever the auditor has
  not supplied one, and clearly marks those ratings as "auto-suggested".
* Lets the audit team record their own recommendations, management action
  plans, owners, dates and status inside the app, then export the register.
* Produces an executive summary, charts, department comparisons, an action
  tracker, and a Word committee report.

Run locally:   streamlit run app.py
"""

from __future__ import annotations

import datetime as dt
import hashlib
import html
import io
import re
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

APP_DIR = Path(__file__).parent
SAMPLE_FILE = APP_DIR / "data" / "observations.xlsx"

# ════════════════════════════════════════════════════════════════════════════
# 1. REFERENCE DATA
# ════════════════════════════════════════════════════════════════════════════

RATINGS = ["Critical", "High", "Medium", "Low"]
RATING_COLORS = {"Critical": "#8B1E3F", "High": "#C0392B", "Medium": "#E0A230", "Low": "#4A8C6F"}
RATING_RANK = {r: i for i, r in enumerate(RATINGS)}

STATUSES = ["Awaiting management response", "Open", "In progress", "Implemented", "Closed", "Risk accepted"]
CLOSED_STATUSES = {"Implemented", "Closed", "Risk accepted"}
STATUS_COLORS = {
    "Awaiting management response": "#8A94A6", "Open": "#C0392B", "In progress": "#E0A230",
    "Implemented": "#4A8C6F", "Closed": "#2E5E4E", "Risk accepted": "#5B6C8F",
}

INK = "#1C2733"
NAVY = "#22385C"
MUTED = "#6B7684"
GRID = "#E3E7EC"
# On-screen colours must read on both light and dark themes.
ACCENT = "#4A6FA5"                  # mid blue: visible on white and on near-black
SOFT_GRID = "rgba(128,128,128,0.22)"
SOFT_TEXT = "rgba(128,128,128,0.95)"

COSO_ORDER = ["Control environment", "Risk assessment", "Control activities",
              "Information & communication", "Monitoring activities"]

# ── Template / canonical fields ─────────────────────────────────────────────
# key, header written to Excel, column width, required?, guidance, why it matters
FIELDS = [
    ("obs_id", "Sr.", 6, "Yes", "Serial number or reference (e.g. 1, OBS-01).", "Lets the Committee refer to a finding unambiguously."),
    ("assignment", "Assignment", 18, "Yes", "Audit assignment / engagement name.", "Enables comparison across audits."),
    ("entity", "Entity", 14, "Optional", "Legal entity audited (e.g. KRN, PFSL).", "Group-level view where several entities are audited."),
    ("department", "Department", 18, "Yes", "Function that owns the finding (Finance, HR, IT, Credit...).", "Drives department comparisons and accountability."),
    ("process", "Process / Sub-process", 20, "Optional", "Process audited (e.g. Payroll, Loan disbursement).", "Pinpoints where the control sits."),
    ("title", "Observation Title", 30, "Recommended", "One-line headline of the issue.", "Used in the executive summary and report; auto-generated if blank."),
    ("finding", "Audit Finding", 60, "Yes", "Condition: what was found.", "The core of the observation."),
    ("criteria", "Criteria", 40, "Recommended", "Policy, law, regulation or standard breached.", "Shows the finding is grounded in an obligation."),
    ("cause", "Cause", 40, "Recommended", "Root cause of the gap.", "Recommendations should address causes, not symptoms."),
    ("risk_text", "Risk / Impact", 40, "Recommended", "Consequence if not addressed.", "Explains why the Committee should care."),
    ("risk_rating", "Risk Rating", 12, "Recommended", "Critical / High / Medium / Low.", "Auditor rating overrides the dashboard's auto-suggestion."),
    ("likelihood", "Likelihood (1-5)", 10, "Optional", "1 = rare, 5 = almost certain.", "Feeds the risk heat map precisely."),
    ("impact_score", "Impact (1-5)", 10, "Optional", "1 = minor, 5 = severe.", "Feeds the risk heat map precisely."),
    ("finding_type", "Finding Type", 14, "Optional", "Design / Operating.", "Design = control missing; Operating = control exists but not followed."),
    ("repeat", "Repeat Finding", 10, "Optional", "Yes / No.", "Repeat findings signal weak management follow-through."),
    ("financial_impact", "Financial Impact (PKR)", 14, "Optional", "Quantified exposure, if any.", "Allows the Committee to prioritise by value at risk."),
    ("control_area", "Control Area", 24, "Optional", "Leave blank to auto-classify, or choose from the Lists sheet.", "Auditor choice overrides auto-classification."),
    ("recommendation", "Recommendation", 50, "Yes", "Your recommendation to fix the issue.", "The auditor's call to action."),
    ("mgmt_response", "Management Response", 40, "Recommended", "Management's comments / agreement.", "Shows whether management accepts the finding."),
    ("action_plan", "Management Action Plan", 45, "Yes", "Specific steps management will take.", "What the Committee will track to closure."),
    ("owner", "Action Owner", 18, "Yes", "Named person or role responsible.", "Accountability for implementation."),
    ("target_date", "Timeline", 14, "Yes", "Target date (e.g. 31-Dec-2026).", "Drives overdue tracking."),
    ("revised_target", "Revised Timeline", 14, "Optional", "Revised date if extended.", "Tracks slippage transparently."),
    ("status", "Status", 14, "Yes", "Open / In progress / Implemented / Closed / Risk accepted.", "Tracks implementation progress."),
    ("date_raised", "Date Raised", 12, "Optional", "Date the finding was reported.", "Enables ageing analysis."),
    ("closure_date", "Closure Date", 12, "Optional", "Date the action was verified closed.", "Evidence of follow-up."),
    ("evidence_ref", "Evidence Ref", 14, "Optional", "Working paper / annexure reference.", "Audit trail to supporting evidence."),
    ("auditor", "Auditor", 14, "Optional", "Auditor responsible.", "Internal accountability within the audit team."),
]
FIELD_KEYS = [f[0] for f in FIELDS]
FIELD_HEADER = {f[0]: f[1] for f in FIELDS}

OBS_SYNONYMS = {
    "obs_id": ["sr", "sr no", "s no", "sno", "no", "ref", "ref no", "reference", "obs id", "observation id",
               "observation no", "finding no", "finding id", "id", "serial no", "number", "s r"],
    "assignment": ["assignment", "engagement", "audit assignment", "audit engagement", "audit", "audit name"],
    "entity": ["entity", "company", "subsidiary", "legal entity"],
    "department": ["department", "dept", "function", "business unit", "unit", "division", "department function"],
    "process": ["process", "sub process", "process sub process", "area", "auditable area", "process area"],
    "title": ["observation title", "title", "heading", "finding title", "subject", "short title"],
    "finding": ["audit finding", "finding", "findings", "observation", "observations", "condition", "issue",
                "finding description", "observation details", "audit observation", "audit findings", "details"],
    "criteria": ["criteria", "criterion", "standard", "requirement", "benchmark"],
    "cause": ["cause", "root cause", "causes", "reason"],
    "risk_text": ["risk", "risks", "risk impact", "impact", "effect", "consequence", "risk implication",
                  "implication", "risk exposure", "risk and impact"],
    "risk_rating": ["risk rating", "rating", "severity", "priority", "risk level", "risk category", "criticality"],
    "likelihood": ["likelihood", "likelihood 1 5", "probability"],
    "impact_score": ["impact 1 5", "impact score", "impact rating", "severity score"],
    "finding_type": ["finding type", "deficiency type", "type"],
    "repeat": ["repeat finding", "repeat", "recurring", "repeated", "repeat observation"],
    "financial_impact": ["financial impact", "financial impact pkr", "amount", "amount pkr", "financial exposure",
                         "value at risk", "exposure amount"],
    "control_area": ["control area", "control category", "category", "internal control area"],
    "recommendation": ["recommendation", "recommendations", "auditor recommendation", "audit recommendation",
                       "suggested action"],
    "mgmt_response": ["management response", "management comments", "management comment", "auditee response",
                      "response"],
    "action_plan": ["management action plan", "action plan", "map", "agreed action", "agreed action plan",
                    "corrective action"],
    "owner": ["action owner", "owner", "responsible", "responsible person", "responsibility",
              "person responsible", "responsible officer"],
    "target_date": ["timeline", "target date", "due date", "implementation date", "deadline", "target",
                    "timelines", "implementation timeline"],
    "revised_target": ["revised timeline", "revised target date", "revised date", "extended date"],
    "status": ["status", "implementation status", "action status"],
    "date_raised": ["date raised", "date identified", "date reported", "report date", "date"],
    "closure_date": ["closure date", "date closed", "closed on"],
    "evidence_ref": ["evidence ref", "evidence", "working paper ref", "wp ref", "annexure"],
    "auditor": ["auditor", "prepared by", "audit lead"],
}
QUERY_SYNONYMS = {
    "q_id": ["s no", "sno", "sr", "sr no", "no", "ref", "id"],
    "q_topic": ["points", "point", "topic", "subject", "area", "title", "heading"],
    "q_text": ["description", "query", "queries", "question", "query description", "details"],
    "q_response": ["response", "reply", "management reply", "notes", "remarks", "comments", "auditee response", "answer"],
    "q_status": ["status"],
    "q_owner": ["addressed to", "owner", "responsible"],
}
SKIP_SHEETS = re.compile(r"(instruction|^lists?$|summary|analysis|lookup|readme|guide)", re.I)

# ── Control areas (keyword → weight). Keywords are regex fragments; a leading
#    word boundary is added automatically, so "polic" matches policy/policies.
AREAS = [
    dict(name="Governance & board oversight", coso="Control environment", impact=4,
         kw={r"board(?! of revenue)": 2, r"steering committee": 3, r"oversight": 2, r"governance": 2,
             r"directorship": 3, r"directors?\b": 2, r"audit committee": 3, r"not reported": 2, r"business case": 1.5},
         rec="Define clear oversight and reporting lines to the Board or relevant committee, set a reporting "
             "calendar, and minute and track decisions to closure."),
    dict(name="Organisation structure & HR", coso="Control environment", impact=3,
         kw={r"organogram": 3, r"organi[sz]ation(al)? structure": 3, r"structure of (the )?organi[sz]ation": 3,
             r"vacan": 3, r"positions?\b": 1, r"human resource": 3, r"hr\b": 2, r"head of": 1,
             r"job descriptions?": 2, r"recruit": 2, r"headcount": 2, r"staff": 1},
         rec="Align actual staffing with the Board-approved organogram, fill critical vacancies (including heads "
             "of function) on a time-bound plan, and obtain approval for any deviations."),
    dict(name="Policies & procedures", coso="Control activities", impact=3,
         kw={r"polic": 2, r"procedures?": 2, r"sops?\b": 2, r"manuals?\b": 1, r"framework": 1,
             r"guidelines?": 1.5, r"review frequency": 2, r"approved polic": 2},
         rec="Develop or update the policy/procedure covering roles, responsibilities, approval authority and "
             "review frequency; obtain Board approval and communicate it to relevant staff."),
    dict(name="Regulatory & licensing compliance", coso="Control activities", impact=5,
         kw={r"licen[cs]e": 4, r"renewal": 2, r"regulat": 2, r"secp\b": 3, r"sbp\b": 3, r"nbfc": 3,
             r"statutory": 2, r"non-?compliance": 2, r"compliance": 1.5},
         rec="Maintain a regulatory compliance calendar with named owners, resolve outstanding regulatory matters "
             "with the regulator on priority, and report status to the Board until closed."),
    dict(name="Tax compliance", coso="Control activities", impact=4,
         kw={r"tax": 3, r"fbr\b": 3, r"federal board of revenue": 3, r"withholding": 3, r"tax returns?": 2,
             r"deduct": 1, r"filed": 1},
         rec="Implement a statutory tax calendar with maker-checker review of deduction, deposit and filing; "
             "reconcile deductions to deposits monthly and assess exposure to default surcharge and penalties."),
    dict(name="Credit risk management", coso="Risk assessment", impact=4,
         kw={r"credit": 2.5, r"loans?\b": 2, r"lending": 2, r"concentration": 2.5, r"due dil+igence": 2,
             r"restructur": 2, r"write[- ]?offs?": 2.5, r"portfolio": 1.5, r"borrowers?": 2, r"mark[- ]?up": 2,
             r"investment cases?": 2, r"exposure": 1.5, r"risk participat": 2, r"npls?\b": 2, r"disburs": 2,
             r"collateral": 2},
         rec="Enforce credit policy limits through pre-approval checks, document due-diligence steps with a named "
             "preparer and reviewer, and report exceptions and breaches to the Board Risk Committee."),
    dict(name="Enterprise risk management", coso="Risk assessment", impact=4,
         kw={r"risk regis": 4, r"risk appetite": 4, r"risk assessment": 3, r"risk management": 2,
             r"key risk indicators?": 3, r"kris?\b": 3},
         rec="Maintain risk registers for each function, obtain annual Board approval of the risk appetite "
             "statement, and monitor key risk indicators against it."),
    dict(name="IT access & security", coso="Control activities", impact=4,
         kw={r"user access": 3, r"access matrix": 3, r"roles?\b": 1.5, r"user ids?\b": 2, r"oracle ids?": 3,
             r"deactivat": 3, r"passwords?": 3, r"privileged": 3, r"super users?": 2, r"it logs?": 3,
             r"logs?\b": 2, r"audit trail": 3, r"segregation of duties": 3, r"access\b": 1.5},
         rec="Define a role-based access matrix aligned to job descriptions, perform periodic access reviews, "
             "disable leavers promptly, and retain and review system logs."),
    dict(name="IT operations & continuity", coso="Control activities", impact=3,
         kw={r"back-?ups?": 3, r"disaster recovery": 3, r"business continuity": 3, r"bcp\b": 3,
             r"test scripts?": 2.5, r"uat\b": 2, r"erp\b": 1.5, r"oracle": 1, r"modules?\b": 1.5,
             r"configuration": 1.5, r"integration": 1.5, r"system": 1, r"user manuals?": 2},
         rec="Document and test backup and recovery arrangements, maintain user manuals and test scripts, and "
             "formalise change management for system configurations."),
    dict(name="Financial reporting & period close", coso="Control activities", impact=4,
         kw={r"journal": 2.5, r"jvs?\b": 2.5, r"reopen": 3, r"closed periods?": 3, r"periods?\b": 1,
             r"reconciliation": 3, r"balances?\b": 1.5, r"financial statements?": 2, r"general ledger": 2.5,
             r"gl\b": 2, r"cfo\b": 1.5},
         rec="Introduce an approval hierarchy for journal postings, restrict period reopening to documented CFO "
             "approval, and review system reports and reconciliations monthly."),
    dict(name="Project & contract management", coso="Monitoring activities", impact=3,
         kw={r"contracts?\b": 2.5, r"agreements?": 2, r"vendors?": 1.5, r"milestones?": 3, r"deliverables?": 2.5,
             r"implementation": 1.5, r"business case": 2, r"payments?\b": 1.5, r"sign-?offs?": 2,
             r"projects?\b": 2},
         rec="Ensure contracts include all schedules and deliverables, link payments to signed-off milestones, and "
             "report project progress to a steering committee and the Board."),
    dict(name="Whistleblowing, complaints & ethics", coso="Information & communication", impact=3,
         kw={r"whistle": 4, r"complaints?": 3, r"grievance": 3, r"code of conduct": 3, r"ethic": 2,
             r"conflict of interest": 3, r"website": 1},
         rec="Publish accessible complaint and whistleblowing channels under a Board-approved policy with "
             "confidentiality protections, and report cases periodically to the Audit Committee."),
    dict(name="Custody of assets & records", coso="Control activities", impact=3,
         kw={r"custodian": 3, r"custody": 3, r"safekeeping": 3, r"fixed assets?": 3, r"records?\b": 1.5,
             r"documentation": 1.5, r"frc\b": 1},
         rec="Assign documented custodians for key assets, licences and records, maintain a register, and verify "
             "it periodically."),
    dict(name="Payroll & employee benefits", coso="Control activities", impact=3,
         kw={r"provident fund": 4, r"pf\b": 3, r"gratuity": 3, r"salar": 2, r"payroll": 3, r"eobi": 3,
             r"employee benefits?": 3},
         rec="Establish and maintain employee benefit arrangements (e.g. a separate provident fund) in line with "
             "applicable law and trust deeds, with periodic reconciliation."),
]
OTHER_AREA = dict(name="Other / unclassified", coso="Control activities", impact=3, kw={},
                  rec="Define the control requirement, assign an owner and a target date, and report implementation "
                      "status to the Audit Committee.")
for _a in AREAS:
    _a["_rx"] = [(re.compile(r"\b" + k, re.I), w) for k, w in _a["kw"].items()]
AREA_BY_NAME = {a["name"]: a for a in AREAS + [OTHER_AREA]}
AREA_NAMES = [a["name"] for a in AREAS] + [OTHER_AREA["name"]]

# ── Cross-cutting themes (root-cause patterns) ─────────────────────────────
THEMES = {
    "Absent or deficient policy / procedure": [
        r"\bno (approved |formal |documented )?(polic|procedure)", r"\bpolic\w* (was|were|is|are) not",
        r"not part of the polic", r"fails? to (mention|define|specify)", r"\bno \w+ policy",
        r"approv\w* of \w+ polic", r"there (was|were|is|are) no \w* ?(polic|procedure|restructuring)",
        r"\bno approved polic"],
    "Approval & authorisation gaps": [
        r"not approved", r"without (proper |any )?(approval|authori[sz]ation)", r"approval hierarchy",
        r"\bno approval", r"approval of", r"unauthori[sz]ed"],
    "Delays & missed deadlines": [
        r"delay", r"\blapse", r"within the prescribed", r"due dates?", r"not submitted", r"overdue",
        r"non-completion", r"still in process", r"not (been )?(filed|paid|completed)"],
    "Records, logs & registers not maintained": [
        r"(not|no)\b[\w\s]{0,30}\bm[ai]{1,2}nt[ai]{1,2}\w*", r"\bno (it )?logs", r"no documentation",
        r"not documented", r"no signed", r"not appended", r"no data"],
    "Limit & threshold breaches": [r"breach", r"exceed", r"concentration limit", r"beyond (the )?limit"],
    "Regulatory & statutory exposure": [
        r"income tax", r"\bfbr\b", r"federal board of revenue", r"licen[cs]e", r"regulat", r"statutory",
        r"\bsecp\b", r"penalt", r"rules,? \d{4}", r"act,? \d{4}"],
    "Governance & oversight gaps": [r"not (been )?(reported|presented|approved) (to|by) the board",
                                    r"without (the )?board", r"oversight", r"steering committee", r"not reported",
                                    r"directorship", r"conflict of interest"],
    "Access & segregation of duties": [r"segregation", r"\baccess\b", r"identical roles", r"reopen",
                                       r"deactivat", r"super user"],
    "Capacity & resourcing gaps": [r"vacan", r"\bno head\b", r"there (was|is) no head", r"shortage",
                                   r"understaff"],
    "Practice inconsistent with approved framework": [
        r"inconsisten", r"however,? in actual", r"as against", r"contrary to", r"not in line", r"deviat"],
}
THEMES_RX = {t: [re.compile(p, re.I) for p in pats] for t, pats in THEMES.items()}

DESIGN_RX = [re.compile(p, re.I) for p in [
    r"\bno\s+(approved\s+|formal\s+|documented\s+|separate\s+|seperate\s+)?(polic|procedure|process|mechanism|"
    r"matrix|hierarchy|system|register|framework|head|backup|log)",
    r"there (was|were|is|are) no\b", r"fails? to (mention|define|specify|cover)", r"not part of",
    r"does not (mention|define|specify|cover|include)", r"\babsence of\b", r"not (been )?defined",
    r"\black of\b", r"\ballows?\b", r"^\s*no\b"]]
OPERATING_RX = [re.compile(p, re.I) for p in [
    r"breach", r"delay", r"not (been )?(submitted|paid|filed|deactivated|followed|complied|approved on)",
    r"\blapse", r"inconsisten", r"as against", r"however,? in actual", r"exceed",
    r"not (being |been )?mai?nt\w*", r"still in process", r"\berrors?\b", r"without (proper )?(approval|authori)",
    r"reopen"]]
SEVERITY_RX = [re.compile(p, re.I) for p in [
    r"penalt", r"\bfines?\b", r"fraud|misappropriat|embezzl", r"licen[cs]e", r"breach",
    r"major risk|significant risk|\bmaterial(ly)?\b", r"write[- ]?off", r"\bloss(es)?\b", r"non-?compliance"]]
PERVASIVE_RX = [re.compile(p, re.I) for p in [
    r"\b(all|several|multiple|numerous|various)\b",
    r"\b\d{2,}\s+(former\s+)?(employees|users|cases|clients|instances|transactions|accounts)"]]


# ════════════════════════════════════════════════════════════════════════════
# 2. PARSING
# ════════════════════════════════════════════════════════════════════════════

def norm(s) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(s).lower())).strip()


def clean(v) -> str:
    """Convert any cell value to tidy text ('' for blanks)."""
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, (pd.Timestamp, dt.datetime, dt.date)):
        return v.strftime("%d-%b-%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).replace("\r", "\n").strip()
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return s


_ALL_SYN = {norm(x) for d in (OBS_SYNONYMS, QUERY_SYNONYMS) for v in d.values() for x in v}


def detect_header(raw: pd.DataFrame) -> tuple[int | None, dict]:
    """Find the header row and harvest 'KEY: value' metadata above it."""
    best, best_score = None, 1
    for i in range(min(len(raw), 25)):
        score = sum(1 for v in raw.iloc[i].tolist() if clean(v) and norm(v) in _ALL_SYN)
        if score > best_score:
            best, best_score = i, score
    meta = {}
    if best is not None:
        for i in range(best):
            for v in raw.iloc[i].tolist():
                txt = clean(v)
                if not txt:
                    continue
                m = re.match(r"^\s*([A-Za-z][A-Za-z &/]{1,30}?)\s*:\s*(.*)$", txt)
                if m:
                    key, val = norm(m.group(1)), m.group(2).strip()
                    for canon, keys in {"assignment": ["assignment", "engagement", "audit"],
                                        "period": ["period", "audit period", "review period"],
                                        "department": ["department", "function"],
                                        "entity": ["entity", "company"]}.items():
                        if key in keys and val:
                            meta[canon] = val
                elif "org" not in meta and i == 0:
                    meta["org"] = txt
    return best, meta


def map_columns(columns, synonyms: dict) -> dict:
    rev = {}
    for canon, syns in synonyms.items():
        for s in syns:
            rev.setdefault(norm(s), canon)
    mapping, used = {}, set()
    for c in columns:
        n = norm(c)
        if n in rev and rev[n] not in used:
            mapping[c] = rev[n]
            used.add(rev[n])
    for c in columns:  # fuzzy second pass
        if c in mapping:
            continue
        n = norm(c)
        for canon, syns in synonyms.items():
            if canon in used:
                continue
            if any(len(norm(s)) >= 6 and norm(s) in n for s in syns):
                mapping[c] = canon
                used.add(canon)
                break
    return mapping


def _frame_from_raw(raw: pd.DataFrame, hdr: int) -> pd.DataFrame:
    cols, seen = [], {}
    for j, v in enumerate(raw.iloc[hdr].tolist()):
        name = clean(v) or f"Unnamed {j + 1}"
        seen[name] = seen.get(name, 0) + 1
        cols.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    body = raw.iloc[hdr + 1:].copy()
    body.columns = cols
    return body.reset_index(drop=True)


@st.cache_data(show_spinner=False)
def load_sources(sources: tuple) -> tuple[pd.DataFrame, pd.DataFrame, list, list]:
    """sources = ((filename, bytes), ...). Returns raw observations, queries, metadata, notes."""
    obs_frames, q_frames, metas, notes = [], [], [], []
    for fname, data in sources:
        try:
            if fname.lower().endswith((".csv", ".txt")):
                sheets = {"CSV": pd.read_csv(io.BytesIO(data), header=None, dtype=object)}
            else:
                sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object)
        except Exception as e:  # noqa: BLE001
            notes.append(f"Could not read {fname}: {e}")
            continue
        for sheet, raw in sheets.items():
            if SKIP_SHEETS.search(str(sheet)) or raw.dropna(how="all").empty:
                continue
            hdr, meta = detect_header(raw)
            if hdr is None:
                notes.append(f"{fname} › {sheet}: no recognisable header row, skipped.")
                continue
            body = _frame_from_raw(raw, hdr)
            src = f"{fname} › {sheet}"
            omap = map_columns(body.columns, OBS_SYNONYMS)
            if "finding" in omap.values() or "title" in omap.values() and "recommendation" in omap.values():
                df = body.rename(columns=omap)[list(omap.values())].copy()
                for k in FIELD_KEYS:
                    df[k] = df[k].map(clean) if k in df else ""
                df = df[(df["finding"] != "") | (df["title"] != "")].copy()
                if df.empty:
                    continue
                for k in ("assignment", "department", "entity"):
                    if meta.get(k):
                        df.loc[df[k] == "", k] = meta[k]
                df["period"] = meta.get("period", "")
                df["org"] = meta.get("org", "")
                df["_source"] = src
                df["_uid"] = [hashlib.md5(f"{src}|{i}".encode()).hexdigest()[:10] for i in df.index]
                blank_id = df["obs_id"] == ""
                df.loc[blank_id, "obs_id"] = [str(i + 1) for i in np.where(blank_id)[0]]
                obs_frames.append(df[FIELD_KEYS + ["period", "org", "_source", "_uid"]])
                unmapped = [c for c in body.columns if c not in omap and body[c].map(clean).ne("").any()]
                if unmapped:
                    notes.append(f"{src}: columns not recognised and ignored: {', '.join(map(str, unmapped))}.")
                meta["source"] = src
                metas.append(meta)
                continue
            qmap = map_columns(body.columns, QUERY_SYNONYMS)
            if "q_text" in qmap.values() or "q_topic" in qmap.values():
                if "q_response" not in qmap.values():
                    for c in body.columns:
                        if c not in qmap and str(c).startswith("Unnamed") and body[c].map(clean).ne("").any():
                            qmap[c] = "q_response"
                            break
                q = body.rename(columns=qmap)[list(qmap.values())].copy()
                for k in QUERY_SYNONYMS:
                    q[k] = q[k].map(clean) if k in q else ""
                q = q[(q["q_text"] != "") | (q["q_topic"] != "")].copy()
                q["assignment"] = meta.get("assignment", "")
                q["_source"] = src
                q_frames.append(q[list(QUERY_SYNONYMS) + ["assignment", "_source"]])
            else:
                notes.append(f"{src}: no observation or query columns found, skipped.")
    obs = pd.concat(obs_frames, ignore_index=True) if obs_frames else pd.DataFrame(
        columns=FIELD_KEYS + ["period", "org", "_source", "_uid"])
    qs = pd.concat(q_frames, ignore_index=True) if q_frames else pd.DataFrame(
        columns=list(QUERY_SYNONYMS) + ["assignment", "_source"])
    # Fill assignment across sheets of the same workbook (e.g. Queries sheet with no metadata).
    wb_assign = {m["source"].split(" › ")[0]: m.get("assignment", "") for m in metas if m.get("assignment")}
    if not qs.empty:
        blank = qs["assignment"] == ""
        qs.loc[blank, "assignment"] = qs.loc[blank, "_source"].str.split(" › ").str[0].map(wb_assign).fillna("")
    return obs, qs, metas, notes


# ════════════════════════════════════════════════════════════════════════════
# 3. CLASSIFICATION & SCORING
# ════════════════════════════════════════════════════════════════════════════

def classify_areas(text: str):
    scores = {}
    for a in AREAS:
        s = sum(w for rx, w in a["_rx"] if rx.search(text))
        if s > 0:
            scores[a["name"]] = s
    if not scores:
        return OTHER_AREA["name"], []
    ranked = sorted(scores.items(), key=lambda x: -x[1])
    top = ranked[0][1]
    return ranked[0][0], [n for n, s in ranked[1:] if s >= 0.6 * top][:2]


def detect_themes(text: str) -> list[str]:
    return [t for t, rxs in THEMES_RX.items() if any(rx.search(text) for rx in rxs)]


def normalise_rating(v: str) -> str:
    s = str(v).lower().strip()
    if not s:
        return ""
    if s.startswith("crit") or s in {"very high", "extreme", "1"}:
        return "Critical"
    if s.startswith("high") or s.startswith("significant"):
        return "High"
    if s.startswith("med") or s.startswith("moderate"):
        return "Medium"
    if s.startswith("low") or s.startswith("minor"):
        return "Low"
    return ""


def normalise_status(v: str) -> str:
    s = str(v).lower().strip()
    if not s:
        return ""
    if "accept" in s:
        return "Risk accepted"
    if "closed" in s:
        return "Closed"
    if any(k in s for k in ("implemented", "complete", "resolved", "done")):
        return "Implemented"
    if any(k in s for k in ("progress", "ongoing", "partial", "wip", "under")):
        return "In progress"
    if "await" in s or "no response" in s:
        return "Awaiting management response"
    return "Open"


def parse_date(v: str):
    s = str(v).strip()
    if not s:
        return pd.NaT
    m = re.match(r"^q([1-4])[\s\-/]*(\d{4})$", s, re.I) or re.match(r"^(\d{4})[\s\-/]*q([1-4])$", s, re.I)
    if m:
        a, b = m.groups()
        q, y = (int(a), int(b)) if len(a) == 1 else (int(b), int(a))
        return pd.Timestamp(year=y, month=q * 3, day=1) + pd.offsets.MonthEnd(0)
    with np.errstate(all="ignore"):
        try:
            ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
        except Exception:  # noqa: BLE001
            ts = pd.NaT
    if pd.notna(ts) and re.fullmatch(r"[A-Za-z]+[\s\-,]*\d{4}", s):  # "June 2026" → month end
        ts = ts + pd.offsets.MonthEnd(0)
    return ts


def short_title(text: str, n: int = 95) -> str:
    if not text:
        return ""
    first = text.strip().split("\n")[0]
    m = list(re.finditer(r"\b(?:noted|observed|found)\s+that\s+", first, re.I))
    if m:
        t = first[m[-1].end():]
    else:
        m2 = re.match(r"^(?:while|during)\s+(?:the\s+)?(?:reviewing|review(?:\s+of)?|audit(?:\s+of)?)\s+(?:the\s+)?"
                      r"(.+?)(?:,|\s+following\b|\s+it\s+was\b|:|$)", first, re.I)
        t = f"Anomalies noted in {m2.group(1)}" if m2 and re.search(r"anomal|following|:", first, re.I) else first
    t = re.split(r"(?<=[.?!])\s+", t.strip())[0].strip().rstrip(":;, ")
    if t.endswith(".") and not t.endswith("..."):
        t = t[:-1]
    if len(t) > n:
        t = t[:n].rsplit(" ", 1)[0].rstrip(",;:") + "…"
    return t[:1].upper() + t[1:]


def rating_from_score(score: float) -> str:
    if score >= 20:
        return "Critical"
    if score >= 12:
        return "High"
    if score >= 6:
        return "Medium"
    return "Low"


def _to_int(v, lo=1, hi=5):
    try:
        x = int(float(str(v).strip()))
        return x if lo <= x <= hi else None
    except (TypeError, ValueError):
        return None


def enrich(raw: pd.DataFrame, default_dept: str, as_at: dt.date) -> pd.DataFrame:
    """Pure function: raw canonical fields → full analytical frame."""
    if raw.empty:
        return raw.assign(**{c: pd.Series(dtype=object) for c in [
            "ref", "title_final", "area", "rating", "score", "status_final", "department_final"]})
    d = raw.copy()
    today = pd.Timestamp(as_at)
    d["department_final"] = d["department"].where(d["department"] != "", default_dept or "Unassigned")
    d["assignment_final"] = d["assignment"].where(d["assignment"] != "", "General")
    d["entity_final"] = d["entity"].where(d["entity"] != "", "Not specified")
    d["process_final"] = d["process"].where(d["process"] != "", "Not specified")
    d["title_final"] = [t if t else short_title(f) for t, f in zip(d["title"], d["finding"])]

    # References: numeric → OBS-01; prefix with assignment if refs collide across sources.
    refs = [f"OBS-{int(x):02d}" if re.fullmatch(r"\d+", x) else x for x in d["obs_id"]]
    d["ref"] = refs
    if d["ref"].duplicated(keep=False).any():
        pre = d["assignment_final"].map(lambda s: re.sub(r"[^A-Z]", "", s.upper())[:4] or "AUD")
        d["ref"] = pre + "-" + d["ref"]
    if d["ref"].duplicated(keep=False).any():
        d["ref"] = d["ref"] + "." + d.groupby("ref").cumcount().add(1).astype(str)

    areas, secondaries, cosos, themes, ftypes, Ls, Is, scores, autos, finals, rsrc, asrc, flags = ([] for _ in range(13))
    for _, r in d.iterrows():
        body = " ".join([r["title"], r["finding"], r["criteria"], r["cause"]])
        # Area
        auto_area, sec = classify_areas(body)
        user_area = next((a for a in AREA_NAMES if norm(a) == norm(r["control_area"])), "")
        if r["control_area"] and not user_area:
            user_area = r["control_area"]
        area = user_area or auto_area
        areas.append(area)
        asrc.append("Auditor" if user_area else "Auto")
        secondaries.append("; ".join(s for s in sec if s != area))
        cosos.append(AREA_BY_NAME.get(area, OTHER_AREA)["coso"])
        themes.append(detect_themes(r["finding"] + " " + r["cause"] + " " + r["title"]))
        # Finding type
        ft_user = r["finding_type"].lower()
        if "design" in ft_user and "operat" in ft_user:
            ft = "Design & operating"
        elif "design" in ft_user:
            ft = "Design"
        elif "operat" in ft_user or "effective" in ft_user:
            ft = "Operating"
        else:
            nd = sum(bool(rx.search(r["finding"])) for rx in DESIGN_RX)
            no = sum(bool(rx.search(r["finding"])) for rx in OPERATING_RX)
            ft = ("Design" if nd > no else "Operating" if no > nd else
                  "Design & operating" if nd else "Not determined")
        ftypes.append(ft)
        # Likelihood × impact
        text_all = body + " " + r["risk_text"]
        imp = AREA_BY_NAME.get(area, OTHER_AREA)["impact"] + (1 if any(rx.search(text_all) for rx in SEVERITY_RX) else 0)
        words = len(r["finding"].split())
        is_query = words < 6 or (r["finding"].strip().endswith("?") and words < 15)
        pervasive = (len(re.findall(r"(?m)^\s*\d+[.)]\s", r["finding"])) >= 3 or
                     any(rx.search(r["finding"]) for rx in PERVASIVE_RX))
        repeat = str(r["repeat"]).strip().lower() in {"yes", "y", "true", "1", "repeat"}
        lik = 3 + (ft in ("Operating", "Design & operating")) + repeat + pervasive - is_query
        L = _to_int(r["likelihood"]) or int(np.clip(lik, 1, 5))
        I = _to_int(r["impact_score"]) or int(np.clip(imp, 1, 5))
        Ls.append(L)
        Is.append(I)
        scores.append(L * I)
        auto = rating_from_score(L * I)
        autos.append(auto)
        user_rating = normalise_rating(r["risk_rating"]) or normalise_rating(r["risk_text"].split()[0] if r["risk_text"] else "")
        finals.append(user_rating or auto)
        rsrc.append("Auditor" if user_rating else "Auto-suggested")
        # Quality flags
        f = []
        if is_query:
            f.append("Finding reads as a query or note; elaborate before reporting")
        for k, lbl in [("criteria", "criteria"), ("cause", "cause"), ("risk_text", "risk / impact"),
                       ("recommendation", "recommendation"), ("action_plan", "management action plan"),
                       ("owner", "action owner"), ("target_date", "target date")]:
            if not r[k]:
                f.append(f"No {lbl}")
        if not user_rating:
            f.append("Risk rating auto-suggested; validate")
        flags.append(f)

    d["area"], d["area_source"], d["secondary_areas"], d["coso"] = areas, asrc, secondaries, cosos
    d["theme_list"] = themes
    d["themes"] = ["; ".join(t) for t in themes]
    d["finding_type_final"] = ftypes
    d["likelihood_final"], d["impact_final"], d["score"] = Ls, Is, scores
    d["auto_rating"], d["rating"], d["rating_source"] = autos, finals, rsrc
    d["rating_rank"] = d["rating"].map(RATING_RANK)
    d["suggested_rec"] = d["area"].map(lambda a: AREA_BY_NAME.get(a, OTHER_AREA)["rec"])
    d["repeat_flag"] = d["repeat"].str.lower().isin({"yes", "y", "true", "1", "repeat"})
    d["financial_value"] = pd.to_numeric(d["financial_impact"].str.replace(r"[^\d.\-]", "", regex=True),
                                         errors="coerce")

    # Status & dates
    st_user = d["status"].map(normalise_status)
    has_map = (d["action_plan"] != "") | (d["mgmt_response"] != "")
    d["status_final"] = np.where(st_user != "", st_user,
                                 np.where(has_map, "Open", "Awaiting management response"))
    tgt_text = d["revised_target"].where(d["revised_target"] != "", d["target_date"])
    d["target_dt"] = tgt_text.map(parse_date)
    d["target_label"] = [t.strftime("%d-%b-%Y") if pd.notna(t) else (x or "Not set")
                         for t, x in zip(d["target_dt"], tgt_text)]
    d["raised_dt"] = d["date_raised"].map(parse_date)
    d["days_to_target"] = (d["target_dt"] - today).dt.days
    open_mask = ~d["status_final"].isin(CLOSED_STATUSES)
    d["is_open"] = open_mask
    d["overdue"] = open_mask & d["days_to_target"].lt(0)

    def bucket(row):
        if row["status_final"] in CLOSED_STATUSES:
            return "Closed / implemented"
        x = row["days_to_target"]
        if pd.isna(x):
            return "No target date"
        if x < -90:
            return "Overdue > 90 days"
        if x < -30:
            return "Overdue 31–90 days"
        if x < 0:
            return "Overdue ≤ 30 days"
        if x <= 30:
            return "Due within 30 days"
        if x <= 90:
            return "Due in 31–90 days"
        return "Due later"
    d["due_bucket"] = d.apply(bucket, axis=1)
    d["age_days"] = (today - d["raised_dt"]).dt.days

    comp_fields = ["finding", "criteria", "cause", "risk_text", "risk_rating", "recommendation",
                   "action_plan", "owner", "target_date", "status", "department"]
    d["completeness"] = (d[comp_fields] != "").sum(axis=1) / len(comp_fields)
    d["flag_list"] = flags
    d["flags"] = ["; ".join(f) for f in flags]
    return d


def classify_queries(q: pd.DataFrame) -> pd.DataFrame:
    if q.empty:
        return q.assign(area=pd.Series(dtype=object), q_state=pd.Series(dtype=object))
    q = q.copy()
    q["area"] = [classify_areas(f"{a} {b}")[0] for a, b in zip(q["q_topic"], q["q_text"])]
    st_norm = q["q_status"].str.lower()
    q["q_state"] = np.where(st_norm.str.contains("clos|resolv|answer|respond"), "Responded",
                            np.where(q["q_response"] != "", "Response received", "Awaiting response"))
    q["ref"] = [f"Q-{int(x):02d}" if re.fullmatch(r"\d+", str(x)) else (x or f"Q-{i + 1:02d}")
                for i, x in enumerate(q["q_id"])]
    return q


# ════════════════════════════════════════════════════════════════════════════
# 4. EXECUTIVE SUMMARY
# ════════════════════════════════════════════════════════════════════════════

def pct(a, b):
    return f"{(100 * a / b):.0f}%" if b else "0%"


def plural(n, word, pl=None):
    return f"{n} {word if n == 1 else (pl or word + 's')}"


def build_summary(d: pd.DataFrame, q: pd.DataFrame) -> dict:
    n = len(d)
    if n == 0:
        return {"headline": "No observations match the current selection.", "paras": [], "matters": [], "caveats": []}
    rc = d["rating"].value_counts()
    crit, high = int(rc.get("Critical", 0)), int(rc.get("High", 0))
    areas = d["area"].value_counts()
    coso = d["coso"].value_counts()
    depts = d["department_final"].nunique()
    assigns = sorted(d["assignment_final"].unique())
    th = pd.Series([t for ts in d["theme_list"] for t in ts]).value_counts()
    ft = d["finding_type_final"].value_counts()
    awaiting = int((d["status_final"] == "Awaiting management response").sum())
    with_map = int((d["action_plan"] != "").sum())
    overdue = int(d["overdue"].sum())
    closed = int(d["status_final"].isin(CLOSED_STATUSES).sum())
    auto_n = int((d["rating_source"] == "Auto-suggested").sum())
    rec_n = int((d["recommendation"] != "").sum())

    headline = (f"{plural(n, 'observation')} reported; {crit + high} "
                f"({pct(crit + high, n)}) rated Critical or High.")
    paras = []
    paras.append(
        f"This summary covers {plural(n, 'observation')} arising from "
        f"{', '.join(assigns) if len(assigns) <= 3 else plural(len(assigns), 'assignment')} across "
        f"{plural(depts, 'department')}. {crit} {'is' if crit == 1 else 'are'} rated Critical, {high} High, "
        f"{int(rc.get('Medium', 0))} Medium and {int(rc.get('Low', 0))} Low. "
        f"The average risk score is {d['score'].mean():.1f} out of 25.")
    top_areas = [f"{a} ({c})" for a, c in areas.head(3).items()]
    paras.append(
        f"Weaknesses are concentrated in {', '.join(top_areas[:-1]) + ' and ' + top_areas[-1] if len(top_areas) > 1 else top_areas[0]}. "
        f"Mapped to the COSO framework, the largest share relates to {coso.index[0].lower()} "
        f"({pct(coso.iloc[0], n)} of findings)"
        + (f", followed by {coso.index[1].lower()} ({pct(coso.iloc[1], n)})." if len(coso) > 1 else "."))
    d_n = int(ft.get("Design", 0) + ft.get("Design & operating", 0))
    o_n = int(ft.get("Operating", 0) + ft.get("Design & operating", 0))
    theme_txt = ""
    if len(th):
        theme_txt = (" The most common cross-cutting themes are " +
                     ", ".join(f"{t.lower()} ({c})" for t, c in th.head(3).items()) + ".")
    paras.append(
        f"{d_n} observations point to design gaps, where a control or policy is absent or incomplete, and "
        f"{o_n} to operating failures, where an existing requirement was not followed.{theme_txt}")
    q_txt = ""
    if not q.empty:
        q_open = int((q["q_state"] == "Awaiting response").sum())
        q_txt = f" In addition, {plural(len(q), 'audit query', 'audit queries')} were raised, of which {q_open} await a response."
    paras.append(
        f"Recommendations have been recorded for {rec_n} of {n} observations and management action plans for "
        f"{with_map}. {awaiting} {'observation awaits' if awaiting == 1 else 'observations await'} a management response, "
        f"{overdue} {'action is' if overdue == 1 else 'actions are'} overdue and {closed} "
        f"{'has' if closed == 1 else 'have'} been implemented or closed.{q_txt}")

    top = d.sort_values(["rating_rank", "score"], ascending=[True, False])
    top = top[top["rating"].isin(["Critical", "High"])].head(6)
    matters = [dict(ref=r["ref"], rating=r["rating"], title=r["title_final"], area=r["area"],
                    dept=r["department_final"], score=r["score"]) for _, r in top.iterrows()]
    caveats = []
    if auto_n:
        caveats.append(f"{plural(auto_n, 'risk rating')} {'was' if auto_n == 1 else 'were'} auto-suggested by the "
                       f"dashboard's likelihood × impact model and should be validated by the audit team before issue.")
    auto_area = int((d["area_source"] == "Auto").sum())
    if auto_area:
        caveats.append(f"Control areas for {auto_area} observations were assigned by keyword classification.")
    return {"headline": headline, "paras": paras, "matters": matters, "caveats": caveats}


# ════════════════════════════════════════════════════════════════════════════
# 5. EXPORTS
# ════════════════════════════════════════════════════════════════════════════

def _style_sheet(ws, widths=None, wrap_cols=None, header_row=1):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="C9D1DB")
    for c in ws[header_row]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    for i, col in enumerate(ws.iter_cols(min_row=header_row, max_row=header_row), start=1):
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = (widths or {}).get(i, 16)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def export_register(raw: pd.DataFrame, d: pd.DataFrame, q: pd.DataFrame) -> bytes:
    """Excel register: 'Observations' is re-uploadable; other sheets hold the analysis."""
    buf = io.BytesIO()
    obs = raw[FIELD_KEYS].rename(columns=FIELD_HEADER)
    analysis = d[["ref", "title_final", "department_final", "assignment_final", "area", "area_source",
                  "secondary_areas", "coso", "themes", "finding_type_final", "likelihood_final", "impact_final",
                  "score", "rating", "rating_source", "status_final", "target_label", "due_bucket",
                  "completeness", "flags", "suggested_rec"]].rename(columns={
        "ref": "Ref", "title_final": "Title", "department_final": "Department", "assignment_final": "Assignment",
        "area": "Control area", "area_source": "Area source", "secondary_areas": "Secondary areas",
        "coso": "COSO component", "themes": "Themes", "finding_type_final": "Design / operating",
        "likelihood_final": "Likelihood", "impact_final": "Impact", "score": "Risk score",
        "rating": "Risk rating", "rating_source": "Rating source", "status_final": "Status",
        "target_label": "Target", "due_bucket": "Due status", "completeness": "Completeness",
        "flags": "Data gaps", "suggested_rec": "Standard recommendation (reference)"})
    by_area = (d.groupby("area").agg(Observations=("ref", "count"),
                                     Critical=("rating", lambda s: (s == "Critical").sum()),
                                     High=("rating", lambda s: (s == "High").sum()),
                                     Medium=("rating", lambda s: (s == "Medium").sum()),
                                     Low=("rating", lambda s: (s == "Low").sum()),
                                     Avg_score=("score", "mean"), Overdue=("overdue", "sum"))
               .sort_values("Observations", ascending=False).reset_index()
               .rename(columns={"area": "Control area", "Avg_score": "Average score"}))
    by_dept = (d.groupby("department_final").agg(Observations=("ref", "count"),
                                                 Critical_High=("rating", lambda s: s.isin(["Critical", "High"]).sum()),
                                                 Avg_score=("score", "mean"),
                                                 Awaiting=("status_final", lambda s: (s == "Awaiting management response").sum()),
                                                 Overdue=("overdue", "sum"))
               .reset_index().rename(columns={"department_final": "Department", "Critical_High": "Critical + High",
                                              "Avg_score": "Average score", "Awaiting": "Awaiting response"}))
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        obs.to_excel(xw, sheet_name="Observations", index=False)
        analysis.to_excel(xw, sheet_name="Analysis", index=False)
        by_area.to_excel(xw, sheet_name="Summary by area", index=False)
        by_dept.to_excel(xw, sheet_name="Summary by department", index=False)
        if not q.empty:
            q[["ref", "q_topic", "q_text", "q_response", "q_state", "area", "assignment"]].rename(columns={
                "ref": "S.No", "q_topic": "Points", "q_text": "Description", "q_response": "Response",
                "q_state": "Status", "area": "Control area", "assignment": "Assignment"}).to_excel(
                xw, sheet_name="Queries", index=False)
        wb = xw.book
        _style_sheet(wb["Observations"], {i + 1: f[2] for i, f in enumerate(FIELDS)})
        _style_sheet(wb["Analysis"], {1: 10, 2: 40, 7: 30, 9: 40, 20: 45, 21: 60})
        for s in ("Summary by area", "Summary by department"):
            _style_sheet(wb[s], {1: 36})
            for row in wb[s].iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0"
        for row in wb["Analysis"].iter_rows(min_row=2, min_col=19, max_col=19):
            for c in row:
                c.number_format = "0%"
        if "Queries" in wb.sheetnames:
            _style_sheet(wb["Queries"], {1: 8, 2: 24, 3: 60, 4: 30, 5: 18, 6: 30})
    return buf.getvalue()


def blank_template(raw: pd.DataFrame | None = None, queries: pd.DataFrame | None = None,
                   title="Internal Audit Department", assignment="", period="") -> bytes:
    """Professional input template with dropdowns, an instructions sheet and a worked example."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Observations"
    ws["A1"] = title.upper() if title else "INTERNAL AUDIT DEPARTMENT"
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color=NAVY.lstrip("#"))
    ws["D1"] = f"ASSIGNMENT: {assignment}"
    ws["G1"] = f"PERIOD: {period}"
    for c in ("D1", "G1"):
        ws[c].font = Font(name="Arial", bold=True, size=10)
    ws.append([])
    ws.append([f[1] for f in FIELDS])
    rows = raw[FIELD_KEYS].values.tolist() if raw is not None and not raw.empty else []
    for r in rows:
        ws.append([_num_or_text(v) for v in r])
    _style_sheet(ws, {i + 1: f[2] for i, f in enumerate(FIELDS)}, header_row=3)
    ws.row_dimensions[3].height = 32
    last = max(ws.max_row, 3) + 300

    lists = wb.create_sheet("Lists")
    cols = {"Risk Rating": RATINGS, "Status": STATUSES[1:], "Finding Type": ["Design", "Operating", "Design & operating"],
            "Repeat Finding": ["Yes", "No"], "Control Area": AREA_NAMES, "Score (1-5)": ["1", "2", "3", "4", "5"]}
    for j, (h, vals) in enumerate(cols.items(), start=1):
        lists.cell(row=1, column=j, value=h)
        for i, v in enumerate(vals, start=2):
            lists.cell(row=i, column=j, value=v)
    _style_sheet(lists, {5: 38})

    def add_dv(field_key, list_col, n):
        col = FIELD_KEYS.index(field_key) + 1
        letter = ws.cell(row=3, column=col).column_letter
        dv = DataValidation(type="list", formula1=f"=Lists!${list_col}$2:${list_col}${n + 1}", allow_blank=True)
        dv.add(f"{letter}4:{letter}{last}")
        ws.add_data_validation(dv)
    add_dv("risk_rating", "A", len(RATINGS))
    add_dv("status", "B", len(STATUSES) - 1)
    add_dv("finding_type", "C", 3)
    add_dv("repeat", "D", 2)
    add_dv("control_area", "E", len(AREA_NAMES))
    add_dv("likelihood", "F", 5)
    add_dv("impact_score", "F", 5)

    qs = wb.create_sheet("Queries")
    qs.append(["S.No", "Points", "Description", "Response", "Status"])
    if queries is not None and not queries.empty:
        for _, r in queries.iterrows():
            qs.append([_num_or_text(r["q_id"]), r["q_topic"], r["q_text"], r["q_response"], r["q_status"]])
    _style_sheet(qs, {1: 7, 2: 26, 3: 70, 4: 36, 5: 16})

    ins = wb.create_sheet("Instructions", 0)
    ins["A1"] = "How to use this template"
    ins["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY.lstrip("#"))
    notes = [
        "Enter one observation per row on the Observations sheet, starting at row 4. Do not rename the header row.",
        "Cells with dropdowns (Risk Rating, Status, Finding Type, Repeat Finding, Control Area, Likelihood, Impact) "
        "take values from the Lists sheet.",
        "Leave Risk Rating or Control Area blank to let the dashboard auto-suggest them; anything you enter overrides it.",
        "Put pending audit queries on the Queries sheet; responses go in the Response column.",
        "Upload this workbook (or several, e.g. one per department) to the dashboard.",
    ]
    for i, t in enumerate(notes, start=3):
        ins.cell(row=i, column=1, value=f"{i - 2}. {t}").font = Font(name="Arial", size=10)
    start = len(notes) + 5
    ins.cell(row=start, column=1, value="Field guide").font = Font(name="Arial", bold=True, size=12)
    hdr = ["Column", "Required", "What to enter", "Why it matters to the Committee", "Worked example"]
    example = {
        "obs_id": "1", "assignment": "Payroll audit FY2026", "entity": "KRN", "department": "Finance",
        "process": "Payroll taxes", "title": "Withholding tax on salaries deposited late",
        "finding": "Tax deducted on salaries for Jan–Mar was deposited with FBR 12–20 days after the statutory deadline.",
        "criteria": "Income Tax Rules, 2002: tax deducted must be deposited within 7 days of the end of each fortnight.",
        "cause": "No tax calendar; deposit depends on one staff member with no reviewer.",
        "risk_text": "Default surcharge and penalties; reputational risk with the tax authority.",
        "risk_rating": "High", "likelihood": "4", "impact_score": "4", "finding_type": "Operating", "repeat": "No",
        "financial_impact": "250000", "control_area": "Tax compliance",
        "recommendation": "Introduce a statutory tax calendar with maker-checker sign-off and monthly reconciliation.",
        "mgmt_response": "Agreed.", "action_plan": "Finance to adopt a tax calendar and assign a reviewer.",
        "owner": "CFO", "target_date": "31-Dec-2026", "revised_target": "", "status": "In progress",
        "date_raised": "15-Sep-2026", "closure_date": "", "evidence_ref": "WP-TX-03", "auditor": "Audit Associate"}
    for j, h in enumerate(hdr, start=1):
        ins.cell(row=start + 1, column=j, value=h)
    for i, f in enumerate(FIELDS, start=start + 2):
        for j, v in enumerate([f[1], f[3], f[4], f[5], example.get(f[0], "")], start=1):
            ins.cell(row=i, column=j, value=v)
    from openpyxl.styles import Border, Side
    for row in ins.iter_rows(min_row=start + 1, max_row=start + 1 + len(FIELDS)):
        for c in row:
            c.font = Font(name="Arial", size=10, bold=c.row == start + 1,
                          color="FFFFFF" if c.row == start + 1 else "000000")
            if c.row == start + 1:
                c.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=Side(style="thin", color="C9D1DB"))
    for letter, w in zip("ABCDE", (24, 12, 48, 48, 48)):
        ins.column_dimensions[letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num_or_text(v):
    s = clean(v)
    return int(s) if re.fullmatch(r"\d{1,6}", s) else (s or None)


def export_docx(d: pd.DataFrame, summary: dict, q: pd.DataFrame, meta: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color.lstrip("#"))
        tcPr.append(shd)

    def table(rows, header, widths=None):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, NAVY)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
                if header[i] == "Rating" and v in RATING_COLORS:
                    shade(cells[i], RATING_COLORS[v])
                    for p in cells[i].paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
        if widths:
            t.autofit = False
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()
        return t

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string(NAVY.lstrip("#"))

    doc.add_heading(meta.get("title") or "Internal audit observations: Audit Committee summary", 0)
    info = doc.add_paragraph()
    for k, v in [("Organisation", meta.get("org")), ("Assignment(s)", ", ".join(sorted(d["assignment_final"].unique()))),
                 ("Period", meta.get("period")), ("Position as at", meta.get("as_at")),
                 ("Prepared by", meta.get("prepared_by")), ("Meeting", meta.get("meeting"))]:
        if v:
            r = info.add_run(f"{k}: ")
            r.bold = True
            info.add_run(f"{v}\n")

    doc.add_heading("1. Executive summary", 1)
    p = doc.add_paragraph()
    p.add_run(summary["headline"]).bold = True
    for para in summary["paras"]:
        doc.add_paragraph(para)
    if summary["matters"]:
        doc.add_heading("Key matters for the Committee's attention", 2)
        for m in summary["matters"]:
            doc.add_paragraph(f"{m['ref']} ({m['rating']}, {m['area']}): {m['title']}", style="List Bullet")

    doc.add_heading("2. Risk profile", 1)
    rc = d["rating"].value_counts()
    table([[r, int(rc.get(r, 0)), pct(int(rc.get(r, 0)), len(d))] for r in RATINGS],
          ["Rating", "Observations", "Share"], [5, 4, 4])

    doc.add_heading("3. Control areas", 1)
    g = d.groupby("area").agg(n=("ref", "count"), ch=("rating", lambda s: s.isin(["Critical", "High"]).sum()),
                              avg=("score", "mean"), coso=("coso", "first")).sort_values("n", ascending=False)
    table([[a, r["coso"], int(r["n"]), int(r["ch"]), f"{r['avg']:.1f}"] for a, r in g.iterrows()],
          ["Control area", "COSO component", "Obs.", "Critical + High", "Avg score"], [5.5, 4.5, 1.6, 2.6, 2.2])

    if d["department_final"].nunique() >= 1:
        doc.add_heading("4. Departments", 1)
        g2 = d.groupby("department_final").agg(
            n=("ref", "count"), ch=("rating", lambda s: s.isin(["Critical", "High"]).sum()),
            aw=("status_final", lambda s: (s == "Awaiting management response").sum()),
            od=("overdue", "sum"), avg=("score", "mean"))
        table([[k, int(r["n"]), int(r["ch"]), int(r["aw"]), int(r["od"]), f"{r['avg']:.1f}"] for k, r in g2.iterrows()],
              ["Department", "Obs.", "Critical + High", "Awaiting response", "Overdue", "Avg score"])

    doc.add_heading("5. Critical and high-rated observations", 1)
    top = d[d["rating"].isin(["Critical", "High"])].sort_values(["rating_rank", "score"], ascending=[True, False])
    if top.empty:
        doc.add_paragraph("No Critical or High observations in this selection.")
    for _, r in top.iterrows():
        doc.add_heading(f"{r['ref']}: {r['title_final']}", 2)
        rows = [("Rating", f"{r['rating']} (score {r['score']}/25, {r['rating_source'].lower()})"),
                ("Department", r["department_final"]), ("Control area", r["area"]),
                ("Finding", r["finding"]), ("Criteria", r["criteria"] or "Not provided"),
                ("Cause", r["cause"] or "Not provided"), ("Risk / impact", r["risk_text"] or "Not provided"),
                ("Recommendation", r["recommendation"] or "To be provided by the audit team"),
                ("Management action plan", r["action_plan"] or "Awaiting management response"),
                ("Owner / target", f"{r['owner'] or 'Not assigned'} / {r['target_label']}"),
                ("Status", r["status_final"])]
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.autofit = False
        for k, v in rows:
            c = t.add_row().cells
            c[0].text, c[1].text = k, str(v)
            shade(c[0], "EEF1F5")
            c[0].paragraphs[0].runs[0].bold = True
            c[0].width, c[1].width = Cm(4), Cm(13)
            for cell in c:
                for pp in cell.paragraphs:
                    for run in pp.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()

    if not q.empty:
        doc.add_heading("6. Open audit queries", 1)
        oq = q[q["q_state"] == "Awaiting response"]
        doc.add_paragraph(f"{len(q)} queries raised; {len(oq)} awaiting a response.")
        if not oq.empty:
            table([[r["ref"], r["q_topic"], r["q_text"]] for _, r in oq.iterrows()],
                  ["Ref", "Topic", "Query"], [1.6, 4, 11.4])

    doc.add_heading("Appendix A: observation register", 1)
    reg = d.sort_values(["rating_rank", "score"], ascending=[True, False])
    table([[r["ref"], r["title_final"], r["area"], r["rating"], r["status_final"], r["target_label"]]
           for _, r in reg.iterrows()], ["Ref", "Observation", "Control area", "Rating", "Status", "Target"],
          [1.8, 6.2, 3.6, 1.8, 2.2, 1.9])

    doc.add_heading("Appendix B: method", 1)
    for t in [
        "Control areas are assigned by the auditor where provided; otherwise by weighted keyword classification of "
        "the finding, criteria and cause. Each area maps to a COSO 2013 internal control component.",
        "Where the auditor has not rated an observation, the dashboard suggests one from a 5 × 5 likelihood × impact "
        "model: impact starts from the control area's inherent impact (+1 for regulatory, breach, loss or fraud "
        "indicators); likelihood starts at 3 (+1 for operating failures, repeat findings or pervasive issues; −1 for "
        "findings that read as queries). Score ≥ 20 Critical, 12–19 High, 6–11 Medium, ≤ 5 Low.",
        *summary["caveats"]]:
        doc.add_paragraph(t, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════════
# 6. CHARTS
# ════════════════════════════════════════════════════════════════════════════

FONT = "Public Sans, Segoe UI, Arial, sans-serif"
PLOT_CFG = {"displaylogo": False, "modeBarButtonsToRemove": ["lasso2d", "select2d"]}


def style(fig, height=360, legend=True):
    fig.update_layout(
        height=height, margin=dict(l=8, r=8, t=78 if legend else 48, b=8),
        font=dict(family=FONT, size=12),
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", showlegend=legend,
        legend=dict(orientation="h", yanchor="bottom", y=1.01, x=0, xanchor="left", title=None, font=dict(size=11)),
        title=dict(font=dict(size=14), x=0.005, xref="container", xanchor="left",
                   y=0.985, yref="container", yanchor="top"),
        hoverlabel=dict(font_family=FONT))
    fig.update_xaxes(gridcolor=SOFT_GRID, zerolinecolor=SOFT_GRID, title_font=dict(size=11))
    fig.update_yaxes(gridcolor=SOFT_GRID, zerolinecolor=SOFT_GRID, title_font=dict(size=11))
    return fig


def show(fig):
    st.plotly_chart(fig, config=PLOT_CFG)


def fig_rating_donut(d):
    c = d["rating"].value_counts().reindex(RATINGS, fill_value=0)
    c = c[c > 0]
    fig = go.Figure(go.Pie(labels=c.index, values=c.values, hole=0.62, sort=False,
                           marker=dict(colors=[RATING_COLORS[r] for r in c.index], line=dict(color="white", width=2)),
                           textinfo="value", hovertemplate="%{label}: %{value} (%{percent})<extra></extra>"))
    fig.add_annotation(text=f"<b>{len(d)}</b><br><span style='font-size:11px'>observations</span>",
                       showarrow=False, font=dict(size=24, family=FONT))
    return style(fig, 320).update_layout(title="Risk rating mix")


def fig_area_stack(d, height=None):
    g = d.groupby(["area", "rating"]).size().reset_index(name="n")
    order = d["area"].value_counts().index.tolist()[::-1]
    fig = px.bar(g, y="area", x="n", color="rating", orientation="h", color_discrete_map=RATING_COLORS,
                 category_orders={"area": order, "rating": RATINGS},
                 labels={"area": "", "n": "Observations", "rating": "Rating"})
    fig.update_traces(hovertemplate="%{y}<br>%{fullData.name}: %{x}<extra></extra>")
    return style(fig, height or max(300, 34 * len(order) + 90)).update_layout(title="Observations by control area",
                                                                            bargap=0.35)


def fig_heatmap(d):
    counts = np.zeros((5, 5), dtype=int)
    refs = [[[] for _ in range(5)] for _ in range(5)]
    for _, r in d.iterrows():
        li, im = int(r["likelihood_final"]) - 1, int(r["impact_final"]) - 1
        counts[li, im] += 1
        refs[li][im].append(r["ref"])
    zone = np.array([[(i + 1) * (j + 1) for j in range(5)] for i in range(5)])
    zone_idx = np.vectorize(lambda s: 3 if s >= 20 else 2 if s >= 12 else 1 if s >= 6 else 0)(zone)
    cs = [[0, "#DCEBE3"], [0.25, "#DCEBE3"], [0.25, "#F7E6C0"], [0.5, "#F7E6C0"],
          [0.5, "#F1C3BC"], [0.75, "#F1C3BC"], [0.75, "#D9A7B6"], [1, "#D9A7B6"]]
    text = [[str(counts[i, j]) if counts[i, j] else "" for j in range(5)] for i in range(5)]
    hover = [[f"Likelihood {i + 1} × Impact {j + 1} = {(i + 1) * (j + 1)}<br>"
              f"{counts[i, j]} observation(s)<br>{', '.join(refs[i][j][:8])}" for j in range(5)] for i in range(5)]
    fig = go.Figure(go.Heatmap(z=zone_idx, x=[1, 2, 3, 4, 5], y=[1, 2, 3, 4, 5], colorscale=cs, zmin=0, zmax=3,
                               showscale=False, text=text, texttemplate="<b>%{text}</b>",
                               textfont=dict(size=16, color=INK), customdata=hover,
                               hovertemplate="%{customdata}<extra></extra>", xgap=3, ygap=3))
    fig.update_xaxes(title="Impact", tickvals=[1, 2, 3, 4, 5], showgrid=False)
    fig.update_yaxes(title="Likelihood", tickvals=[1, 2, 3, 4, 5], showgrid=False)
    return style(fig, 340, legend=False).update_layout(title="Risk heat map (likelihood × impact)")


def fig_coso(d):
    c = d["coso"].value_counts().reindex(COSO_ORDER, fill_value=0)[::-1]
    avg = d.groupby("coso")["score"].mean().reindex(c.index)
    fig = go.Figure(go.Bar(y=c.index, x=c.values, orientation="h", marker_color=ACCENT, text=c.values,
                           textposition="outside", customdata=avg.fillna(0).round(1),
                           hovertemplate="%{y}<br>%{x} observation(s)<br>Avg score %{customdata}<extra></extra>"))
    fig.update_xaxes(title="Observations", range=[0, max(c.max(), 1) * 1.18])
    return style(fig, 340, legend=False).update_layout(title="COSO 2013 component coverage", bargap=0.4)


def fig_treemap(d):
    g = d.groupby(["coso", "area"]).agg(n=("ref", "count"), avg=("score", "mean")).reset_index()
    top = g.groupby("coso").apply(lambda x: pd.Series({"n": x["n"].sum(),
                                                      "avg": (x["avg"] * x["n"]).sum() / x["n"].sum()}),
                                  include_groups=False).reset_index()
    ids = list(top["coso"]) + [f"{c}/{a}" for c, a in zip(g["coso"], g["area"])]
    labels = list(top["coso"]) + list(g["area"])
    parents = [""] * len(top) + list(g["coso"])
    values = list(top["n"]) + list(g["n"])
    avgs = list(top["avg"]) + list(g["avg"])
    fig = go.Figure(go.Treemap(
        ids=ids, labels=labels, parents=parents, values=values, branchvalues="total",
        marker=dict(colors=avgs, colorscale=["#DCEBE3", "#E0A230", "#C0392B", "#8B1E3F"], cmin=4, cmax=22,
                    showscale=True, colorbar=dict(title="Avg<br>score", thickness=10),
                    line=dict(color="rgba(128,128,128,0.35)", width=1.5)),
        customdata=np.round(avgs, 1), texttemplate="<b>%{label}</b><br>%{value} obs.",
        textfont=dict(size=13), tiling=dict(pad=4), pathbar=dict(visible=False), root_color="rgba(0,0,0,0)",
        hovertemplate="%{label}<br>%{value} observation(s)<br>Avg score %{customdata}<extra></extra>"))
    return style(fig, 440, legend=False).update_layout(title="Where weaknesses sit (colour = average risk score)")


def fig_themes(d):
    s = pd.Series([t for ts in d["theme_list"] for t in ts]).value_counts()
    if s.empty:
        return None
    s = s.sort_values()
    fig = go.Figure(go.Bar(y=s.index, x=s.values, orientation="h", marker_color=ACCENT,
                           hovertemplate="%{y}: %{x}<extra></extra>", text=s.values, textposition="outside"))
    return style(fig, max(260, 32 * len(s) + 80), legend=False).update_layout(
        title="Cross-cutting themes (an observation can show several)", xaxis_title="Observations")


def fig_type(d):
    c = d["finding_type_final"].value_counts()
    cols = {"Design": ACCENT, "Operating": "#C0392B", "Design & operating": "#8E5572", "Not determined": "#B5BDC8"}
    fig = go.Figure(go.Pie(labels=c.index, values=c.values, hole=0.55,
                           marker=dict(colors=[cols.get(k, "#999") for k in c.index], line=dict(color="white", width=2))))
    return style(fig, 300).update_layout(title="Design gaps vs operating failures")


def fig_group_rating(d, by, label):
    g = d.groupby([by, "rating"]).size().reset_index(name="n")
    fig = px.bar(g, x=by, y="n", color="rating", barmode="group", color_discrete_map=RATING_COLORS,
                 category_orders={"rating": RATINGS}, labels={by: label, "n": "Observations", "rating": "Rating"})
    return style(fig, 360).update_layout(title=f"Risk ratings by {label.lower()}")


def fig_group_area_heat(d, by, label):
    p = pd.crosstab(d[by], d["area"])
    fig = go.Figure(go.Heatmap(z=p.values, x=p.columns, y=p.index, colorscale=["#DCE4EF", "#8FA6C8", "#2F4F7F"],
                               text=p.values, texttemplate="%{text}", showscale=False, xgap=2, ygap=2,
                               hovertemplate=f"{label}: %{{y}}<br>%{{x}}: %{{z}}<extra></extra>"))
    fig.update_xaxes(tickangle=-35)
    return style(fig, max(280, 60 * len(p.index) + 200), legend=False).update_layout(
        title=f"Control areas by {label.lower()}")


def fig_radar(d, by):
    areas = d["area"].value_counts().head(8).index.tolist()
    fig = go.Figure()
    palette = [ACCENT, "#C0392B", "#4A8C6F", "#E0A230", "#8E5572", "#3B7EA1", "#6D1A36", "#7A8B99"]
    for i, (k, sub) in enumerate(d.groupby(by)):
        vals = [int((sub["area"] == a).sum()) for a in areas]
        fig.add_trace(go.Scatterpolar(r=vals + vals[:1], theta=areas + areas[:1], name=str(k), fill="toself",
                                      opacity=0.55, line=dict(color=palette[i % len(palette)])))
    fig.update_layout(polar=dict(bgcolor="rgba(0,0,0,0)", radialaxis=dict(gridcolor=SOFT_GRID),
                                 angularaxis=dict(gridcolor=SOFT_GRID, tickfont=dict(size=10))))
    return style(fig, 440).update_layout(title="Control-area profile", margin=dict(l=110, r=110, t=90, b=30))


def fig_exposure(d, by, label):
    g = d.groupby(by).agg(score=("score", "sum"), n=("ref", "count")).reset_index().sort_values("score")
    fig = go.Figure(go.Bar(y=g[by], x=g["score"], orientation="h", marker_color="#8E5572",
                           customdata=g["n"], text=g["score"], textposition="outside",
                           hovertemplate="%{y}<br>Total risk score %{x}<br>%{customdata} observations<extra></extra>"))
    return style(fig, max(240, 44 * len(g) + 100), legend=False).update_layout(
        title=f"Risk exposure index by {label.lower()} (sum of risk scores)", xaxis_title="Total score")


def fig_status(d):
    c = d["status_final"].value_counts().reindex(STATUSES).dropna()
    fig = go.Figure(go.Pie(labels=c.index, values=c.values, hole=0.6, sort=False, textinfo="value",
                           marker=dict(colors=[STATUS_COLORS[s] for s in c.index], line=dict(color="white", width=2))))
    return style(fig, 320).update_layout(title="Implementation status")


def fig_due(d):
    order = ["Overdue > 90 days", "Overdue 31–90 days", "Overdue ≤ 30 days", "Due within 30 days",
             "Due in 31–90 days", "Due later", "No target date", "Closed / implemented"]
    cols = ["#6D1A36", "#C0392B", "#E07B5A", "#E0A230", "#C9B458", "#8FA6C8", "#B5BDC8", "#4A8C6F"]
    c = d["due_bucket"].value_counts().reindex(order, fill_value=0)
    fig = go.Figure(go.Bar(x=c.index, y=c.values, marker_color=cols, text=c.values, textposition="outside",
                           hovertemplate="%{x}: %{y}<extra></extra>"))
    fig.update_xaxes(tickangle=-25)
    fig.update_yaxes(range=[0, max(c.max(), 1) * 1.2])
    return style(fig, 340, legend=False).update_layout(title="Action due-date profile", yaxis_title="Observations")


def fig_timeline(d, as_at):
    t = d.dropna(subset=["target_dt"])
    if t.empty:
        return None
    fig = px.scatter(t, x="target_dt", y="area", color="rating", color_discrete_map=RATING_COLORS,
                     category_orders={"rating": RATINGS}, hover_name="ref",
                     hover_data={"title_final": True, "status_final": True, "owner": True, "target_dt": False,
                                 "area": False, "rating": False},
                     labels={"target_dt": "Target date", "area": "", "title_final": "Observation",
                             "status_final": "Status", "owner": "Owner", "rating": "Rating"})
    fig.update_traces(marker=dict(size=13, line=dict(color="white", width=1.5)))
    x0 = pd.Timestamp(as_at)
    fig.add_shape(type="line", x0=x0, x1=x0, y0=0, y1=1, yref="paper", line=dict(color=SOFT_TEXT, dash="dot", width=1))
    fig.add_annotation(x=x0, y=1.02, yref="paper", text="Report date", showarrow=False,
                       font=dict(size=10, color=SOFT_TEXT))
    return style(fig, max(300, 40 * t["area"].nunique() + 120)).update_layout(title="Action timeline")


def fig_owner(d):
    o = d[d["is_open"]].assign(owner_f=lambda x: x["owner"].where(x["owner"] != "", "Not assigned"))
    if o.empty:
        return None
    g = o.groupby(["owner_f", "rating"]).size().reset_index(name="n")
    order = o["owner_f"].value_counts().index.tolist()[::-1]
    fig = px.bar(g, y="owner_f", x="n", color="rating", orientation="h", color_discrete_map=RATING_COLORS,
                 category_orders={"owner_f": order, "rating": RATINGS},
                 labels={"owner_f": "", "n": "Open actions", "rating": "Rating"})
    return style(fig, max(240, 36 * len(order) + 100)).update_layout(title="Open actions by owner")


def fig_completeness(d):
    fields = [("finding", "Finding"), ("criteria", "Criteria"), ("cause", "Cause"), ("risk_text", "Risk / impact"),
              ("risk_rating", "Risk rating (auditor)"), ("recommendation", "Recommendation"),
              ("mgmt_response", "Management response"), ("action_plan", "Action plan"), ("owner", "Owner"),
              ("target_date", "Target date"), ("status", "Status"), ("department", "Department")]
    vals = [(lbl, (d[k] != "").mean() * 100) for k, lbl in fields]
    s = pd.DataFrame(vals, columns=["Field", "pct"]).iloc[::-1]
    colors = ["#4A8C6F" if v >= 90 else "#E0A230" if v >= 50 else "#C0392B" for v in s["pct"]]
    fig = go.Figure(go.Bar(y=s["Field"], x=s["pct"], orientation="h", marker_color=colors,
                           text=[f"{v:.0f}%" for v in s["pct"]], textposition="outside",
                           hovertemplate="%{y}: %{x:.0f}% complete<extra></extra>"))
    fig.update_xaxes(range=[0, 112], ticksuffix="%")
    return style(fig, 420, legend=False).update_layout(title="Field completeness across observations")


# ════════════════════════════════════════════════════════════════════════════
# 7. UI HELPERS
# ════════════════════════════════════════════════════════════════════════════

CSS = f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;500;600;700&family=Source+Serif+4:opsz,wght@8..60,400;8..60,600&display=swap');
/* Colours are deliberately NOT set on text: everything inherits from the active Streamlit theme,
   so the dashboard reads correctly in both light and dark mode. Only borders and tints use
   semi-transparent greys, which work on any background. */
html, body, .stApp {{ font-family: 'Public Sans', 'Segoe UI', Arial, sans-serif; }}
.stApp h1, .stApp h2, .stApp h3 {{ font-family: 'Source Serif 4', Georgia, serif; letter-spacing: -0.01em; }}
.stApp h1 {{ font-weight: 600; font-size: 2.05rem; }}
.block-container {{ padding-top: 2.2rem; max-width: 1400px; }}
.muted, .subtle, .caveat, .kpi-note, .field-label, .matter small {{ opacity: 0.72; }}
.subtle {{ font-size: 0.92rem; margin-top: -0.6rem; margin-bottom: 1.1rem; }}
.kpi {{ border: 1px solid {SOFT_GRID}; border-top: 4px solid var(--accent, {ACCENT}); border-radius: 6px;
        padding: 0.8rem 0.95rem 0.7rem; background: rgba(128,128,128,0.06); height: 100%; }}
.kpi-value {{ font-size: 1.9rem; font-weight: 700; line-height: 1.1; font-variant-numeric: tabular-nums; }}
.kpi-label {{ font-size: 0.86rem; font-weight: 600; margin-top: 0.25rem; }}
.kpi-note {{ font-size: 0.78rem; margin-top: 0.15rem; }}
.strip {{ display: flex; height: 16px; border-radius: 3px; overflow: hidden; margin: 0.2rem 0 0.35rem; }}
.strip div {{ height: 100%; }}
.strip-legend {{ display: flex; gap: 1.2rem; flex-wrap: wrap; font-size: 0.84rem; margin-bottom: 1.2rem; }}
.strip-legend span.dot {{ display: inline-block; width: 10px; height: 10px; border-radius: 2px; margin-right: 6px; vertical-align: -1px; }}
.brief {{ font-family: 'Source Serif 4', Georgia, serif; font-size: 1.04rem; line-height: 1.62;
          border-left: 3px solid {ACCENT}; padding: 0.2rem 0 0.2rem 1.1rem; max-width: 78ch; }}
.brief p {{ margin: 0 0 0.7rem; }}
.brief .lead {{ font-weight: 600; font-size: 1.12rem; }}
.caveat {{ font-size: 0.82rem; max-width: 90ch; margin-bottom: 0.7rem; }}
.pill {{ display: inline-block; padding: 2px 9px; border-radius: 999px; font-size: 0.78rem; font-weight: 600;
         margin: 0 6px 6px 0; background: rgba(128,128,128,0.18); }}
.pill.r {{ color: #FFFFFF; }}
.field-label {{ font-size: 0.78rem; font-weight: 700; margin: 0.6rem 0 0.1rem; }}
.field-empty {{ opacity: 0.5; font-style: italic; }}
.matter {{ padding: 0.55rem 0.8rem; border: 1px solid {SOFT_GRID}; border-left: 4px solid var(--c); border-radius: 4px;
           margin-bottom: 0.45rem; background: rgba(128,128,128,0.05); }}
.matter b {{ font-variant-numeric: tabular-nums; }}
</style>
"""


def kpi(col, value, label, note="", accent=NAVY):
    col.markdown(f"<div class='kpi' style='--accent:{accent}'><div class='kpi-value'>{value}</div>"
                 f"<div class='kpi-label'>{html.escape(label)}</div><div class='kpi-note'>{html.escape(note)}</div></div>",
                 unsafe_allow_html=True)


def risk_strip(d):
    c = d["rating"].value_counts().reindex(RATINGS, fill_value=0)
    n = max(len(d), 1)
    segs = "".join(f"<div style='width:{100 * v / n:.2f}%;background:{RATING_COLORS[r]}' title='{r}: {v}'></div>"
                   for r, v in c.items() if v)
    leg = "".join(f"<span><span class='dot' style='background:{RATING_COLORS[r]}'></span>{r} <b>{v}</b> "
                  f"<span class='muted'>({pct(v, n)})</span></span>" for r, v in c.items())
    st.markdown(f"<div class='strip'>{segs}</div><div class='strip-legend'>{leg}</div>", unsafe_allow_html=True)


def page_header(title, sub):
    st.markdown(f"# {title}")
    st.markdown(f"<div class='subtle'>{html.escape(sub)}</div>", unsafe_allow_html=True)


def pill(text, color=None):
    if color:
        return f"<span class='pill r' style='background:{color}'>{html.escape(text)}</span>"
    return f"<span class='pill'>{html.escape(text)}</span>"


def field(label, value, empty="Not yet provided"):
    v = html.escape(value).replace("\n", "<br>") if value else f"<span class='field-empty'>{empty}</span>"
    return f"<div class='field-label'>{label}</div><div>{v}</div>"


def summary_html(s):
    body = f"<p class='lead'>{html.escape(s['headline'])}</p>" + "".join(f"<p>{html.escape(p)}</p>" for p in s["paras"])
    return f"<div class='brief'>{body}</div>"


def matters_html(matters):
    return "".join(
        f"<div class='matter' style='--c:{RATING_COLORS[m['rating']]}'><b>{m['ref']}</b> "
        f"{pill(m['rating'], RATING_COLORS[m['rating']])} {html.escape(m['title'])}<br>"
        f"<small>{html.escape(m['area'])} | {html.escape(m['dept'])} | score {m['score']}/25</small></div>"
        for m in matters)


# ════════════════════════════════════════════════════════════════════════════
# 8. PAGES
# ════════════════════════════════════════════════════════════════════════════

def page_overview(d, q, ctx):
    page_header("Audit observations overview", ctx["subtitle"])
    if d.empty:
        st.info("No observations match the current filters. Clear a filter in the sidebar to see results.")
        return
    risk_strip(d)
    n = len(d)
    ch = int(d["rating"].isin(["Critical", "High"]).sum())
    aw = int((d["status_final"] == "Awaiting management response").sum())
    od = int(d["overdue"].sum())
    rec = int((d["recommendation"] != "").sum())
    auto = int((d["rating_source"] == "Auto-suggested").sum())
    c = st.columns(6)
    kpi(c[0], n, "Observations", f"{d['department_final'].nunique()} dept., {d['area'].nunique()} control areas")
    kpi(c[1], ch, "Critical + high", pct(ch, n) + " of total", RATING_COLORS["Critical"])
    kpi(c[2], f"{d['score'].mean():.1f}", "Average risk score", "out of 25", "#8E5572")
    kpi(c[3], aw, "Awaiting management", "no response or action plan yet", STATUS_COLORS["Awaiting management response"])
    kpi(c[4], od, "Overdue actions", "past target date, still open", RATING_COLORS["High"] if od else "#4A8C6F")
    kpi(c[5], f"{pct(rec, n)}", "With recommendation", f"{auto} ratings auto-suggested", RATING_COLORS["Medium"])
    st.write("")
    s = build_summary(d, q)
    left, right = st.columns([1.35, 1])
    with left:
        st.markdown("### Executive summary")
        st.markdown(summary_html(s), unsafe_allow_html=True)
        for cv in s["caveats"]:
            st.markdown(f"<div class='caveat'>{html.escape(cv)}</div>", unsafe_allow_html=True)
    with right:
        st.markdown("### Key matters for attention")
        if s["matters"]:
            st.markdown(matters_html(s["matters"]), unsafe_allow_html=True)
        else:
            st.caption("No Critical or High observations in this selection.")
    st.divider()
    a, b = st.columns([1, 1.5])
    with a:
        show(fig_rating_donut(d))
    with b:
        show(fig_area_stack(d, height=330 if d["area"].nunique() < 8 else None))
    a, b = st.columns(2)
    with a:
        show(fig_heatmap(d))
    with b:
        show(fig_coso(d))


def page_areas(d, q, ctx):
    page_header("Internal control areas", "Auto-detected control areas, COSO mapping and cross-cutting themes")
    if d.empty:
        st.info("No observations match the current filters.")
        return
    show(fig_treemap(d))
    a, b = st.columns([1.4, 1])
    with a:
        f = fig_themes(d)
        if f:
            show(f)
    with b:
        show(fig_type(d))
    st.markdown("### Area detail")
    g = (d.groupby("area").agg(Observations=("ref", "count"),
                               Critical=("rating", lambda s: int((s == "Critical").sum())),
                               High=("rating", lambda s: int((s == "High").sum())),
                               Medium=("rating", lambda s: int((s == "Medium").sum())),
                               Low=("rating", lambda s: int((s == "Low").sum())),
                               Avg=("score", "mean"), COSO=("coso", "first"),
                               Departments=("department_final", "nunique"))
         .sort_values(["Observations", "Avg"], ascending=False).reset_index())
    st.dataframe(g, hide_index=True, column_config={
        "area": st.column_config.TextColumn("Control area", width="large"),
        "Avg": st.column_config.ProgressColumn("Avg score", min_value=0, max_value=25, format="%.1f"),
        "COSO": "COSO component"})
    pick = st.selectbox("Look inside a control area", g["area"].tolist())
    sub = d[d["area"] == pick].sort_values(["rating_rank", "score"], ascending=[True, False])
    st.markdown(f"<div class='caveat'>{html.escape(AREA_BY_NAME.get(pick, OTHER_AREA)['coso'])}. "
                f"Standard recommendation for reference: {html.escape(AREA_BY_NAME.get(pick, OTHER_AREA)['rec'])}</div>",
                unsafe_allow_html=True)
    st.dataframe(sub[["ref", "title_final", "rating", "score", "department_final", "finding_type_final", "themes",
                      "secondary_areas"]], hide_index=True, column_config={
        "ref": "Ref", "title_final": st.column_config.TextColumn("Observation", width="large"), "rating": "Rating",
        "score": "Score", "department_final": "Department", "finding_type_final": "Type", "themes": "Themes",
        "secondary_areas": "Also touches"})


def page_departments(d, q, ctx):
    page_header("Department comparison", "Compare risk profiles across departments, assignments, entities or processes")
    if d.empty:
        st.info("No observations match the current filters.")
        return
    options = {"Department": "department_final", "Assignment": "assignment_final", "Entity": "entity_final",
               "Process": "process_final", "Control area": "area"}
    label = st.radio("Compare by", list(options), horizontal=True)
    by = options[label]
    k = d[by].nunique()
    if k == 1:
        st.info(f"All observations in this selection belong to one {label.lower()} ({d[by].iloc[0]}). The profile "
                f"below still applies; side-by-side comparisons appear automatically once you load observations with "
                f"more than one {label.lower()} (add a Department column or upload one file per department).")
    tbl = (d.groupby(by).agg(Observations=("ref", "count"),
                             Critical=("rating", lambda s: int((s == "Critical").sum())),
                             High=("rating", lambda s: int((s == "High").sum())),
                             Avg=("score", "mean"), Exposure=("score", "sum"),
                             Awaiting=("status_final", lambda s: int((s == "Awaiting management response").sum())),
                             Overdue=("overdue", "sum"), Complete=("completeness", "mean"),
                             Top_area=("area", lambda s: s.value_counts().index[0]))
           .sort_values("Exposure", ascending=False).reset_index())
    st.dataframe(tbl, hide_index=True, column_config={
        by: label, "Avg": st.column_config.NumberColumn("Avg score", format="%.1f"),
        "Exposure": st.column_config.NumberColumn("Exposure index", help="Sum of risk scores"),
        "Awaiting": "Awaiting response", "Overdue": st.column_config.NumberColumn("Overdue", format="%d"),
        "Complete": st.column_config.ProgressColumn("Data completeness", min_value=0, max_value=1, format="percent"),
        "Top_area": "Most frequent control area"})
    a, b = st.columns(2)
    with a:
        show(fig_group_rating(d, by, label))
    with b:
        show(fig_exposure(d, by, label))
    if by != "area":
        a, b = st.columns([1.3, 1])
        with a:
            show(fig_group_area_heat(d, by, label))
        with b:
            show(fig_radar(d, by))


def page_actions(d, q, ctx):
    page_header("Action tracker", f"Management action plans, owners and due dates as at {ctx['as_at']:%d %B %Y}")
    if d.empty:
        st.info("No observations match the current filters.")
        return
    c = st.columns(5)
    kpi(c[0], int(d["is_open"].sum()), "Open actions", "not implemented or closed")
    kpi(c[1], int(d["overdue"].sum()), "Overdue", "past target date", RATING_COLORS["Critical"])
    kpi(c[2], int((d["due_bucket"] == "Due within 30 days").sum()), "Due in 30 days", "", RATING_COLORS["Medium"])
    kpi(c[3], int((d["due_bucket"] == "No target date").sum()), "No target date", "needs a timeline", "#8A94A6")
    kpi(c[4], int(d["status_final"].isin(CLOSED_STATUSES).sum()), "Implemented / closed", "", "#4A8C6F")
    st.write("")
    a, b = st.columns([1, 1.4])
    with a:
        show(fig_status(d))
    with b:
        show(fig_due(d))
    f = fig_timeline(d, ctx["as_at"])
    if f:
        show(f)
    else:
        st.info("Target dates have not been recorded yet. Add them in the Timeline column (or on the "
                "Recommendations & actions page) to see the action timeline and overdue analysis.")
    f = fig_owner(d)
    if f:
        show(f)
    st.markdown("### Overdue and upcoming")
    watch = d[d["is_open"] & d["days_to_target"].notna()].sort_values("days_to_target")
    if watch.empty:
        st.caption("Nothing to show until target dates are recorded.")
    else:
        st.dataframe(watch[["ref", "title_final", "rating", "owner", "target_label", "days_to_target", "status_final",
                            "department_final"]], hide_index=True, column_config={
            "ref": "Ref", "title_final": st.column_config.TextColumn("Observation", width="large"), "rating": "Rating",
            "owner": "Owner", "target_label": "Target", "days_to_target": st.column_config.NumberColumn(
                "Days to target", help="Negative = overdue"), "status_final": "Status",
            "department_final": "Department"})


def page_register(d, q, ctx):
    page_header("Observation register", "Full detail for every observation, in the criteria-condition-cause-effect format")
    if d.empty:
        st.info("No observations match the current filters.")
        return
    a, b = st.columns([2, 1])
    sort = a.radio("Sort by", ["Risk (highest first)", "Reference", "Target date"], horizontal=True)
    limit = b.selectbox("Show", [25, 50, 100, "All"], index=0)
    if sort.startswith("Risk"):
        v = d.sort_values(["rating_rank", "score"], ascending=[True, False])
    elif sort == "Reference":
        v = d.sort_values("ref")
    else:
        v = d.sort_values("target_dt", na_position="last")
    if limit != "All":
        v = v.head(int(limit))
    for _, r in v.iterrows():
        label = f"{r['ref']}: {r['title_final']}  ({r['rating']}, {r['status_final'].lower()})"
        with st.expander(label):
            pills = (pill(r["rating"], RATING_COLORS[r["rating"]]) + pill(f"Score {r['score']}/25") +
                     pill(r["area"]) + pill(r["coso"]) + pill(r["finding_type_final"]) + pill(r["department_final"]) +
                     (pill("Repeat finding", "#6D1A36") if r["repeat_flag"] else ""))
            st.markdown(pills, unsafe_allow_html=True)
            if r["rating_source"] == "Auto-suggested":
                st.caption(f"Rating auto-suggested (likelihood {r['likelihood_final']} × impact {r['impact_final']}). "
                           "Set it on the Recommendations & actions page to confirm or override.")
            c1, c2 = st.columns(2)
            c1.markdown(field("Condition (finding)", r["finding"]) + field("Criteria", r["criteria"]) +
                        field("Cause", r["cause"]) + field("Risk / impact", r["risk_text"]), unsafe_allow_html=True)
            c2.markdown(field("Recommendation", r["recommendation"]) +
                        field("Management response", r["mgmt_response"]) +
                        field("Management action plan", r["action_plan"], "Awaiting management response") +
                        field("Owner and target", f"{r['owner'] or 'Not assigned'}, {r['target_label']}"),
                        unsafe_allow_html=True)
            extra = []
            if r["themes"]:
                extra.append(f"Themes: {r['themes']}")
            if r["secondary_areas"]:
                extra.append(f"Also touches: {r['secondary_areas']}")
            if r["flags"]:
                extra.append(f"Data gaps: {r['flags']}")
            if extra:
                st.markdown("<div class='caveat'>" + "<br>".join(html.escape(x) for x in extra) + "</div>",
                            unsafe_allow_html=True)
            if not r["recommendation"]:
                st.markdown(f"<div class='caveat'><b>Standard wording for reference:</b> "
                            f"{html.escape(r['suggested_rec'])}</div>", unsafe_allow_html=True)


def _save_raw(raw, nonce, msg):
    st.session_state["raw"] = raw
    st.session_state["edit_nonce"] = nonce + 1
    st.session_state["flash"] = msg
    st.rerun()


def page_editor(d, q, ctx):
    page_header("Recommendations & actions",
                "Record your recommendations, confirm ratings and capture management action plans")
    if st.session_state.get("flash"):
        st.success(st.session_state.pop("flash"))
    if d.empty:
        st.info("No observations match the current filters.")
        return
    nonce = st.session_state.get("edit_nonce", 0)
    raw = st.session_state["raw"]
    mode = st.radio("Editing mode", ["One observation at a time", "Table (bulk edit)"], horizontal=True,
                    help="One at a time suits writing full recommendations; the table suits quick bulk updates.")
    st.caption("Ratings and control areas you set replace the auto-suggestions. Edits last for this session, so "
               "download the updated register at the bottom of this page and keep it as your master file.")

    if mode.startswith("One"):
        order = d.sort_values("ref")
        uids = order["_uid"].tolist()
        labels = dict(zip(order["_uid"], order["ref"] + ": " + order["title_final"]))
        if st.session_state.get("one_sel") not in uids:
            st.session_state["one_sel"] = uids[0]
        uid = st.selectbox("Observation", uids, format_func=lambda u: labels.get(u, u), key="one_sel")
        r = d[d["_uid"] == uid].iloc[0]
        rr = raw[raw["_uid"] == uid].iloc[0]
        st.markdown(field("Condition (finding)", r["finding"]), unsafe_allow_html=True)
        st.markdown(f"<div class='caveat'>Auto-suggestion: {r['auto_rating']} (likelihood {r['likelihood_final']} × "
                    f"impact {r['impact_final']} = {r['score']}), {html.escape(r['area'])}, "
                    f"{html.escape(r['finding_type_final'].lower())}. Standard wording for this area: "
                    f"{html.escape(r['suggested_rec'])}</div>", unsafe_allow_html=True)
        with st.form(f"one_{uid}_{nonce}", border=False):
            c1, c2, c3 = st.columns(3)
            r_opts = [f"Auto-suggest ({r['auto_rating']})"] + RATINGS
            cur_r = normalise_rating(rr["risk_rating"])
            rating = c1.selectbox("Risk rating", r_opts, index=r_opts.index(cur_r) if cur_r else 0)
            a_opts = ["Auto-classify"] + AREA_NAMES
            cur_a = next((a for a in AREA_NAMES if norm(a) == norm(rr["control_area"])), "")
            area = c2.selectbox("Control area", a_opts, index=a_opts.index(cur_a) if cur_a else 0)
            s_opts = ["Derive from action plan"] + STATUSES
            cur_s = normalise_status(rr["status"])
            status = c3.selectbox("Status", s_opts, index=s_opts.index(cur_s) if cur_s else 0)
            c1, c2, c3 = st.columns(3)
            criteria = c1.text_area("Criteria", rr["criteria"], height=110)
            cause = c2.text_area("Cause", rr["cause"], height=110)
            risk_text = c3.text_area("Risk / impact", rr["risk_text"], height=110)
            recommendation = st.text_area("Your recommendation", rr["recommendation"], height=130)
            c1, c2 = st.columns(2)
            mgmt = c1.text_area("Management response", rr["mgmt_response"], height=100)
            plan = c2.text_area("Management action plan", rr["action_plan"], height=100)
            c1, c2, c3, c4 = st.columns(4)
            owner = c1.text_input("Action owner", rr["owner"])
            target = c2.text_input("Timeline", rr["target_date"], help="e.g. 31-Dec-2026 or Q1 2027")
            dept = c3.text_input("Department", rr["department"], help=f"Blank = {ctx['default_dept']}")
            is_rep = rr["repeat"].lower() in {"yes", "y", "true", "1", "repeat"}
            repeat = c4.selectbox("Repeat finding", ["No", "Yes"], index=1 if is_rep else 0)
            if st.form_submit_button("Save observation", type="primary"):
                updates = dict(risk_rating="" if rating.startswith("Auto") else rating,
                               control_area="" if area == "Auto-classify" else area,
                               status="" if status.startswith("Derive") else status,
                               criteria=criteria.strip(), cause=cause.strip(), risk_text=risk_text.strip(),
                               recommendation=recommendation.strip(), mgmt_response=mgmt.strip(),
                               action_plan=plan.strip(), owner=owner.strip(), target_date=target.strip(),
                               department=dept.strip(), repeat=repeat if (repeat == "Yes" or rr["repeat"]) else "")
                for k, v in updates.items():
                    raw.loc[raw["_uid"] == uid, k] = v
                _save_raw(raw, nonce, f"Saved {r['ref']}.")
    else:
        view = d[["_uid", "ref", "title_final", "area", "rating", "recommendation", "mgmt_response", "action_plan",
                  "owner", "target_date", "status_final", "department"]].copy().reset_index(drop=True)
        with st.form("editor_form", border=False):
            edited = st.data_editor(
                view, key=f"editor_{nonce}", hide_index=True, num_rows="fixed",
                height=min(640, 42 + 36 * len(view)), disabled=["_uid", "ref", "title_final"],
                column_config={
                    "_uid": None,
                    "ref": st.column_config.TextColumn("Ref", width="small"),
                    "title_final": st.column_config.TextColumn("Observation", width="medium"),
                    "area": st.column_config.SelectboxColumn("Control area", options=AREA_NAMES, width="medium"),
                    "rating": st.column_config.SelectboxColumn("Risk rating", options=RATINGS, width="small"),
                    "recommendation": st.column_config.TextColumn("Your recommendation", width="large"),
                    "mgmt_response": st.column_config.TextColumn("Management response", width="medium"),
                    "action_plan": st.column_config.TextColumn("Management action plan", width="large"),
                    "owner": st.column_config.TextColumn("Action owner", width="small"),
                    "target_date": st.column_config.TextColumn("Timeline", help="e.g. 31-Dec-2026, Q1 2027",
                                                               width="small"),
                    "status_final": st.column_config.SelectboxColumn("Status", options=STATUSES, width="medium"),
                    "department": st.column_config.TextColumn("Department", width="small",
                                                              help=f"Blank = {ctx['default_dept']}"),
                })
            saved = st.form_submit_button("Save changes", type="primary")
        if saved:
            field_map = {"area": "control_area", "rating": "risk_rating", "recommendation": "recommendation",
                         "mgmt_response": "mgmt_response", "action_plan": "action_plan", "owner": "owner",
                         "target_date": "target_date", "status_final": "status", "department": "department"}
            changes = 0
            for i in range(len(view)):
                uid = view.at[i, "_uid"]
                for col, target in field_map.items():
                    old, new = clean(view.at[i, col]), clean(edited.at[i, col])
                    if old != new:
                        raw.loc[raw["_uid"] == uid, target] = new
                        changes += 1
            _save_raw(raw, nonce, f"Saved {changes} change{'s' if changes != 1 else ''}.")

    with st.expander("Standard recommendation wording by control area (for reference)"):
        st.caption("Starting points only. Your recommendations should address the specific cause you found.")
        st.dataframe(pd.DataFrame([(a["name"], a["rec"]) for a in AREAS + [OTHER_AREA]],
                                  columns=["Control area", "Standard wording"]), hide_index=True)
        blanks = d[d["recommendation"] == ""]
        if len(blanks) and st.button(f"Copy standard wording into the {len(blanks)} blank recommendations as drafts"):
            for _, r in blanks.iterrows():
                raw.loc[raw["_uid"] == r["_uid"], "recommendation"] = "[Draft] " + r["suggested_rec"]
            _save_raw(raw, nonce, "Draft wording added. Each is marked [Draft] so you can tailor it.")
    st.download_button("Download updated register (.xlsx)", export_register(raw, ctx["all"], q),
                       file_name=f"audit_register_{ctx['as_at']:%Y%m%d}.xlsx", type="primary",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def page_queries(d, q, ctx):
    page_header("Audit queries", "Open questions raised during fieldwork, linked to control areas")
    if q.empty:
        st.info("No queries sheet was found. Add a sheet with columns such as S.No, Points, Description and "
                "Response to track pending audit queries here.")
        return
    c = st.columns(3)
    kpi(c[0], len(q), "Queries raised")
    kpi(c[1], int((q["q_state"] == "Awaiting response").sum()), "Awaiting response", "", RATING_COLORS["High"])
    kpi(c[2], int((q["q_state"] != "Awaiting response").sum()), "Response received", "", "#4A8C6F")
    st.write("")
    g = q.groupby(["area", "q_state"]).size().reset_index(name="n")
    order = q["area"].value_counts().index.tolist()[::-1]
    fig = px.bar(g, y="area", x="n", color="q_state", orientation="h",
                 color_discrete_map={"Awaiting response": "#C0392B", "Response received": "#E0A230",
                                     "Responded": "#4A8C6F"},
                 category_orders={"area": order}, labels={"area": "", "n": "Queries", "q_state": "Status"})
    show(style(fig, max(260, 36 * len(order) + 100)).update_layout(title="Queries by control area"))
    st.dataframe(q[["ref", "q_topic", "q_text", "q_response", "q_state", "area", "assignment"]], hide_index=True,
                 column_config={"ref": "Ref", "q_topic": "Topic",
                                "q_text": st.column_config.TextColumn("Query", width="large"),
                                "q_response": "Response / notes", "q_state": "Status", "area": "Control area",
                                "assignment": "Assignment"})


def page_report(d, q, ctx):
    page_header("Committee report", "Auto-generated summary and downloadable reports for the Audit Committee")
    if d.empty:
        st.info("No observations match the current filters.")
        return
    s = build_summary(d, q)
    a, b = st.columns([1.5, 1])
    with a:
        st.markdown(summary_html(s), unsafe_allow_html=True)
        for cv in s["caveats"]:
            st.markdown(f"<div class='caveat'>{html.escape(cv)}</div>", unsafe_allow_html=True)
    with b:
        st.markdown("### Report details")
        title = st.text_input("Report title", "Internal audit observations: Audit Committee summary")
        prepared = st.text_input("Prepared by", "Internal Audit Department")
        meeting = st.text_input("Committee meeting (optional)", "")
        meta = dict(title=title, prepared_by=prepared, meeting=meeting, as_at=f"{ctx['as_at']:%d %B %Y}",
                    org=ctx["org"], period=ctx["period"])
        st.caption("Reports reflect the filters currently applied in the sidebar.")
        st.download_button("Download committee report (.docx)", export_docx(d, s, q, meta), type="primary",
                           file_name=f"audit_committee_report_{ctx['as_at']:%Y%m%d}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        st.download_button("Download register and analysis (.xlsx)",
                           export_register(st.session_state["raw"][st.session_state["raw"]["_uid"].isin(d["_uid"])], d, q),
                           file_name=f"audit_register_{ctx['as_at']:%Y%m%d}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.download_button("Download analysis (.csv)",
                           d.drop(columns=["theme_list", "flag_list", "_uid"]).to_csv(index=False).encode("utf-8"),
                           file_name="audit_observations_analysis.csv", mime="text/csv")
    st.markdown("### Key matters for attention")
    st.markdown(matters_html(s["matters"]) or "No Critical or High observations.", unsafe_allow_html=True)


def page_quality(d, q, ctx):
    page_header("Data quality & template", "What is missing, and what to add for committee-grade reporting")
    if not d.empty:
        a, b = st.columns([1, 1])
        with a:
            show(fig_completeness(d))
        with b:
            st.markdown("### Readiness")
            ready = d[d["completeness"] >= 0.9]
            kpi(st, f"{pct(len(ready), len(d))}", "Committee-ready observations",
                "≥ 90% of key fields completed", "#4A8C6F" if len(ready) == len(d) else RATING_COLORS["Medium"])
            st.write("")
            flag_counts = pd.Series([f for fl in d["flag_list"] for f in fl]).value_counts()
            st.dataframe(flag_counts.rename_axis("Gap").reset_index(name="Observations"), hide_index=True)
        st.markdown("### Gaps by observation")
        st.dataframe(d.sort_values("completeness")[["ref", "title_final", "completeness", "flags"]], hide_index=True,
                     column_config={"ref": "Ref", "title_final": st.column_config.TextColumn("Observation", width="medium"),
                                    "completeness": st.column_config.ProgressColumn("Complete", min_value=0,
                                                                                     max_value=1, format="percent"),
                                    "flags": st.column_config.TextColumn("What's missing", width="large")})
    if ctx["notes"]:
        with st.expander("File reading notes"):
            for n_ in ctx["notes"]:
                st.write(n_)
    st.markdown("### Recommended columns")
    st.markdown("The dashboard works with whatever columns you have. Adding these makes the Committee pack "
                "complete and the analysis more precise.")
    st.dataframe(pd.DataFrame([(f[1], f[3], f[4], f[5]) for f in FIELDS],
                              columns=["Column", "Required", "What to enter", "Why it matters"]), hide_index=True)
    st.download_button("Download blank input template (.xlsx)", blank_template(),
                       file_name="audit_observations_template.xlsx", type="primary",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with st.expander("How auto-classification and rating work"):
        st.markdown(
            "- **Control area**: weighted keyword matching on the finding, criteria and cause across 14 "
            "control areas, each mapped to a COSO 2013 component. Anything you enter in *Control Area* wins.\n"
            "- **Themes**: pattern detection for recurring root causes (missing policies, delays, approval "
            "gaps, breaches and so on). One observation can carry several themes.\n"
            "- **Design vs operating**: phrases such as *no policy* or *fails to mention* indicate a design gap; "
            "*breached*, *delay* or *not submitted* indicate an operating failure.\n"
            "- **Risk rating**: likelihood × impact on a 5 × 5 grid. Impact starts from the control area's inherent "
            "impact (+1 for regulatory, breach, loss or fraud language); likelihood starts at 3 (+1 operating "
            "failure, +1 repeat, +1 pervasive, −1 if the finding reads as a query). ≥ 20 Critical, 12–19 High, "
            "6–11 Medium, ≤ 5 Low. Your *Risk Rating*, *Likelihood* and *Impact* columns always override.")


# ════════════════════════════════════════════════════════════════════════════
# 9. MAIN
# ════════════════════════════════════════════════════════════════════════════

def password_gate() -> bool:
    try:
        pw = st.secrets.get("APP_PASSWORD", "")
    except Exception:  # noqa: BLE001  (no secrets file)
        pw = ""
    if not pw or st.session_state.get("authed"):
        return True
    st.markdown("# Audit observations dashboard")
    entered = st.text_input("Password", type="password")
    if entered:
        if entered == pw:
            st.session_state["authed"] = True
            st.rerun()
        st.error("Incorrect password.")
    return False


def main():
    st.set_page_config(page_title="Audit Observations Dashboard", page_icon="🛡️", layout="wide",
                       initial_sidebar_state="expanded")
    st.markdown(CSS, unsafe_allow_html=True)
    if not password_gate():
        return

    with st.sidebar:
        st.markdown("### Audit observations")
        uploads = st.file_uploader("Upload observation schedules", type=["xlsx", "xls", "csv"],
                                   accept_multiple_files=True,
                                   help="One or more files. Each sheet is detected as observations or queries.")
        use_sample = False
        if not uploads and SAMPLE_FILE.exists():
            use_sample = st.toggle("Use the file stored in the repository", value=True,
                                   help=f"data/{SAMPLE_FILE.name}")

    sources = tuple((f.name, f.getvalue()) for f in uploads) if uploads else (
        ((SAMPLE_FILE.name, SAMPLE_FILE.read_bytes()),) if use_sample else ())
    if not sources:
        page_header("Audit observations dashboard", "Upload an observation schedule to begin")
        st.markdown("Upload one or more Excel or CSV files in the sidebar. The dashboard finds the header row, "
                    "matches your column names and works with whatever fields are present.")
        st.download_button("Download blank input template (.xlsx)", blank_template(),
                           file_name="audit_observations_template.xlsx", type="primary")
        return

    sig = hashlib.md5(b"".join(n.encode() + hashlib.md5(b).digest() for n, b in sources)).hexdigest()
    if st.session_state.get("sig") != sig:
        raw, queries, metas, notes = load_sources(sources)
        st.session_state.update(sig=sig, raw=raw.copy(), queries=queries, metas=metas, notes=notes, edit_nonce=0)
    raw, queries, metas = st.session_state["raw"], st.session_state["queries"], st.session_state["metas"]

    with st.sidebar:
        default_dept = st.text_input("Department for rows without one", value="Unassigned",
                                     help="Applied wherever the Department column is blank.")
        as_at = st.date_input("Position as at", value=dt.date.today(), format="DD/MM/YYYY",
                              help="Used for overdue and ageing calculations.")
        st.divider()
        page = st.radio("View", ["Overview", "Control areas", "Department comparison", "Action tracker",
                                 "Observation register", "Recommendations & actions", "Audit queries",
                                 "Committee report", "Data quality & template"], label_visibility="collapsed")

    full = enrich(raw, default_dept, as_at)
    q = classify_queries(queries)
    if full.empty:
        page_header("Audit observations dashboard", "No observations found")
        st.warning("No observation rows were found. The sheet needs a header row with a column such as "
                   "'Audit Finding' or 'Observation'.")
        for n_ in st.session_state.get("notes", []):
            st.caption(n_)
        return

    with st.sidebar:
        with st.expander("Filters", expanded=False):
            def ms(label, col, order=None):
                opts = order or sorted(full[col].unique().tolist())
                opts = [o for o in opts if o in set(full[col])]
                return st.multiselect(label, opts, placeholder="All")
            f_dept = ms("Department", "department_final")
            f_asg = ms("Assignment", "assignment_final")
            f_rat = ms("Risk rating", "rating", RATINGS)
            f_area = ms("Control area", "area")
            f_st = ms("Status", "status_final", STATUSES)
            f_txt = st.text_input("Search text", placeholder="e.g. tax, policy, Oracle")
        st.caption(f"{len(full)} observations, {len(q)} queries loaded")

    d = full
    for vals, col in [(f_dept, "department_final"), (f_asg, "assignment_final"), (f_rat, "rating"),
                      (f_area, "area"), (f_st, "status_final")]:
        if vals:
            d = d[d[col].isin(vals)]
    if f_txt:
        blob = (d["title_final"] + " " + d["finding"] + " " + d["recommendation"] + " " + d["cause"]).str.lower()
        d = d[blob.str.contains(re.escape(f_txt.lower()), na=False)]
    qf = q[q["assignment"].isin(f_asg)] if f_asg and not q.empty else q

    meta0 = metas[0] if metas else {}
    orgs = sorted({m.get("org", "") for m in metas if m.get("org")})
    ctx = dict(as_at=as_at, default_dept=default_dept, all=full, notes=st.session_state.get("notes", []),
               org=", ".join(orgs).title() if orgs else "", period=meta0.get("period", ""))
    bits = [ctx["org"] or None, ", ".join(sorted(d["assignment_final"].unique())) or None,
            f"Position as at {as_at:%d %B %Y}"]
    ctx["subtitle"] = " | ".join(b for b in bits if b)

    {"Overview": page_overview, "Control areas": page_areas, "Department comparison": page_departments,
     "Action tracker": page_actions, "Observation register": page_register,
     "Recommendations & actions": page_editor, "Audit queries": page_queries, "Committee report": page_report,
     "Data quality & template": page_quality}[page](d, qf, ctx)


if __name__ == "__main__":
    main()